#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
КОМПЛЕКСНАЯ СИСТЕМА ПАРСИНГА ОБРАЗОВАТЕЛЬНЫХ ОРГАНИЗАЦИЙ
=========================================================
Объединяет 31 модуль сбора данных из различных источников в единый
конвейер обработки. Центральный оркестратор управляет запуском
компонентов, ведёт журнал событий и контролирует версионность
выходных файлов (бэкапы и контрольные точки сохранения).
"""

# ========== ИМПОРТЫ ==========
import requests
from bs4 import BeautifulSoup
import pandas as pd
import time
import os
import re
import urllib3
import signal
import sys
import shutil
import logging
import json
import io
import tempfile
import warnings
from datetime import datetime
from random import uniform
from urllib.parse import urljoin
from pathlib import Path
from functools import reduce
from typing import List, Dict

# Selenium
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.firefox.options import Options as FirefoxOptions

# PDF
import camelot
import pdfplumber
import PyPDF2

# Dotenv
from dotenv import load_dotenv

# Отключаем предупреждения SSL
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
warnings.filterwarnings("ignore", category=urllib3.exceptions.InsecureRequestWarning)

# Загружаем .env для парсеров, требующих авторизации
load_dotenv()

# ========== ОБЩИЕ УТИЛИТЫ ДЛЯ ОРКЕСТРАТОРА ==========

def setup_logging():
    """Настраивает логирование: вывод в консоль и в файл."""
    log_formatter = logging.Formatter(
        '%(asctime)s - %(levelname)s - %(message)s',
        datefmt='%Y-%m-%d %H:%M:%S'
    )
    root_logger = logging.getLogger()
    root_logger.setLevel(logging.INFO)
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(log_formatter)
    root_logger.addHandler(console)
    file_handler = logging.FileHandler('orchestrator.log', encoding='utf-8')
    file_handler.setFormatter(log_formatter)
    root_logger.addHandler(file_handler)
    return logging.getLogger(__name__)

logger = setup_logging()

def get_timestamp():
    """Возвращает строку с текущей датой и временем (без пробелов)."""
    return datetime.now().strftime('%Y-%m-%d_%H-%M-%S')

def get_date_short():
    """Возвращает короткую дату (ГГГГ-ММ-ДД)."""
    return datetime.now().strftime('%Y-%m-%d')

def ensure_dir(directory):
    """Создаёт папку, если её нет."""
    if not os.path.exists(directory):
        os.makedirs(directory)
        logger.info(f"Создана папка: {directory}")

def make_backup(file_path, suffix):
    """Копирует file_path в папку backups/ с добавлением suffix."""
    if not os.path.isfile(file_path):
        logger.warning(f"Файл {file_path} не найден, бэкап пропущен.")
        return None
    ensure_dir('backups')
    base = os.path.basename(file_path)
    name, ext = os.path.splitext(base)
    backup_name = f"{name}_{suffix}{ext}"
    backup_path = os.path.join('backups', backup_name)
    try:
        shutil.copy2(file_path, backup_path)
        logger.info(f"Бэкап создан: {backup_path}")
        return backup_path
    except Exception as e:
        logger.error(f"Ошибка бэкапа {backup_path}: {e}")
        return None

def save_snapshot(file_path, source_name):
    """Создаёт копию файла в папке results/ с именем source_name_дата.xlsx."""
    if not os.path.isfile(file_path):
        logger.warning(f"Файл {file_path} не найден, снапшот не создан.")
        return
    ensure_dir('results')
    date_str = get_date_short()
    base = os.path.basename(file_path)
    name, ext = os.path.splitext(base)
    snapshot_name = f"{source_name}_{date_str}{ext}"
    snapshot_path = os.path.join('results', snapshot_name)
    try:
        shutil.copy2(file_path, snapshot_path)
        logger.info(f"Снапшот сохранён: {snapshot_path}")
    except Exception as e:
        logger.error(f"Ошибка создания снапшота: {e}")

# ========== ПАРСЕР №1: АКАДЕМИК ==========

def run_parser_akademik():
    """Парсер с сайта akademik-pro.ru."""
    OUTPUT_FILENAME = 'akademik_data.xlsx'
    URLS = [
        ("https://akademik-pro.ru/nashi_kursy/kursy_povysheniya_kvalifikacii/", "Курс"),
        ("https://akademik-pro.ru/nashi_kursy/professionalnyie_standarty.html", "Стандарт")
    ]
    HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'}

    def parse_table(soup, category):
        table = soup.find('table', class_='main_table')
        if not table:
            return []
        rows = table.find_all('tr')
        if len(rows) < 2:
            return []
        data = []
        for row in rows[1:]:
            cells = row.find_all(['td', 'th'])
            cell_texts = [cell.get_text(strip=True) for cell in cells if cell.get_text(strip=True)]
            if not cell_texts:
                continue
            num_cols = len(cell_texts)
            if num_cols == 3:
                num, name, hours = cell_texts[0], cell_texts[1], cell_texts[2]
                duties = ''
            elif num_cols == 4:
                num, name, duties, hours = cell_texts[0], cell_texts[1], cell_texts[2], cell_texts[3]
            else:
                continue
            num = num.rstrip('.')
            hours = hours.replace('—', '').strip()
            data.append([num, name, duties, hours, category])
        return data

    print("=" * 60)
    print("ПАРСЕР АКАДЕМИК (akademik-pro.ru)")
    print("=" * 60)
    all_data = []
    for url, category in URLS:
        print(f"Загрузка {url} (категория: {category})")
        try:
            response = requests.get(url, headers=HEADERS, timeout=15, verify=False)
            response.raise_for_status()
            soup = BeautifulSoup(response.text, 'html.parser')
            page_data = parse_table(soup, category)
            if page_data:
                all_data.extend(page_data)
                print(f"  Найдено {len(page_data)} записей")
            else:
                print("  Данных не найдено")
        except Exception as e:
            print(f"  Ошибка: {e}")
    if not all_data:
        print("Нет данных для сохранения")
        return
    df = pd.DataFrame(all_data, columns=['№ п/п', 'Наименование программы', 'Должности', 'Кол-во часов', 'Категория'])
    df.to_excel(OUTPUT_FILENAME, index=False)
    print(f"Сохранено {len(df)} записей в {OUTPUT_FILENAME}")

# ========== ПАРСЕР №2: ЦПС ==========

def run_parser_cps():
    """Парсер с сайта cpspec.ru."""
    BASE_URL = "https://cpspec.ru/course-category/page/{}/"
    HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
    PAGES_TO_PARSE = 105
    DELAY = 1
    OUTPUT_FILENAME = 'cps_data.xlsx'

    def extract_categories(block_soup):
        categories_text = block_soup.find(string=re.compile(r'Категории:'))
        if categories_text:
            parent = categories_text.find_parent()
            links = parent.find_all('a', href=True) if parent else []
            return ' / '.join([link.get_text(strip=True) for link in links])
        return ''

    def extract_audience(block_soup):
        audience_div = block_soup.select_one('.tutor-course-details-widget-var.custom__course__widget')
        if audience_div:
            key_span = audience_div.find('span', string=re.compile(r'Аудитория:'))
            if key_span:
                value_span = audience_div.find('span', class_='tutor-course-details-widget-list')
                if value_span:
                    return value_span.get_text(strip=True)
        return ''

    def extract_hours(block_soup):
        hours = []
        hour_buttons = block_soup.select('.variable-item.button-variable-item')
        for btn in hour_buttons:
            title = btn.get('data-title', '')
            if title:
                hours.append(title)
        return ', '.join(hours) if hours else ''

    def extract_prices(block_soup):
        price_block = block_soup.select_one('.flex.variation__wrap__price')
        if price_block:
            old_tag = price_block.find('del')
            new_tag = price_block.find('ins')
            if old_tag and new_tag:
                old_price = re.sub(r'[^\d\s₽]', '', old_tag.get_text(strip=True))
                new_price = re.sub(r'[^\d\s₽]', '', new_tag.get_text(strip=True))
                return old_price, new_price
            else:
                amount = price_block.select_one('.woocommerce-Price-amount')
                if amount:
                    price = re.sub(r'[^\d\s₽]', '', amount.get_text(strip=True))
                    return '', price
        var_price_block = block_soup.select_one('.woocommerce-variation-price .price')
        if var_price_block:
            amount = var_price_block.select_one('.woocommerce-Price-amount')
            if amount:
                price = re.sub(r'[^\d\s₽]', '', amount.get_text(strip=True))
                return '', price
        return '', ''

    def parse_page(page_num):
        url = BASE_URL.format(page_num)
        print(f"Обработка: {url}")
        try:
            resp = requests.get(url, headers=HEADERS, timeout=15, verify=False)
            resp.raise_for_status()
            soup = BeautifulSoup(resp.text, 'html.parser')
            course_blocks = soup.find_all('div', class_='total__custom')
            if not course_blocks:
                print(f"  Блоки не найдены на странице {page_num}")
                return []
            page_data = []
            for block in course_blocks:
                title_tag = block.find('h3')
                if title_tag and title_tag.find('a'):
                    title = title_tag.find('a').get_text(strip=True)
                    link = title_tag.find('a')['href']
                else:
                    title = link = ''
                page_data.append({
                    'Название курса': title,
                    'Ссылка': link,
                    'Категории': extract_categories(block),
                    'Аудитория': extract_audience(block),
                    'Количество часов (варианты)': extract_hours(block),
                    'Старая цена': extract_prices(block)[0],
                    'Цена со скидкой': extract_prices(block)[1],
                    'Страница': page_num
                })
            print(f"  Собрано {len(page_data)} курсов")
            return page_data
        except Exception as e:
            print(f"  Ошибка: {e}")
            return []

    print("=" * 60)
    print("ПАРСЕР ЦПС (cpspec.ru)")
    print("=" * 60)
    all_courses = []
    for page in range(1, PAGES_TO_PARSE + 1):
        data = parse_page(page)
        all_courses.extend(data)
        time.sleep(DELAY)
    if not all_courses:
        print("Нет данных")
        return
    df = pd.DataFrame(all_courses)
    df.to_excel(OUTPUT_FILENAME, index=False)
    print(f"Сохранено {len(all_courses)} курсов в {OUTPUT_FILENAME}")

# ========== ПАРСЕР №3: ГУБКИНСКИЙ УНИВЕРСИТЕТ ==========

def run_parser_gubkin():
    """Парсер с сайта dpo.gubkin.ru (Selenium)."""
    URL = "https://dpo.gubkin.ru/napravleniya-obucheniya"
    OUTPUT_FILENAME = 'gubkin_data.xlsx'
    DELAY_BETWEEN_FILTERS = 3
    WAIT_TIMEOUT = 10
    last_direction_data = []

    def signal_handler(sig, frame):
        print("\nПрерывание. Сохраняем собранные данные...")
        if last_direction_data:
            save_to_excel(last_direction_data, OUTPUT_FILENAME)
        sys.exit(0)

    def save_to_excel(data, filename):
        if not data:
            return
        if os.path.exists(filename):
            try:
                existing_df = pd.read_excel(filename)
                existing_data = existing_df.to_dict('records')
            except Exception as e:
                print(f"Не удалось загрузить существующий файл: {e}")
                existing_data = []
        else:
            existing_data = []
        combined = existing_data + data
        df = pd.DataFrame(combined)
        df = df.drop_duplicates(subset=['Направление', 'Название программы', 'Стоимость'], keep='first')
        cols_order = ['Направление', 'Название программы', 'Ссылка на программу', 'Форматы',
                      'Вид обучения', 'Дата начала', 'Длительность', 'Стоимость']
        df = df[cols_order]
        df = df.sort_values('Направление')
        df.to_excel(filename, index=False)
        print(f"Сохранено {len(df)} записей в {filename}")

    def parse_card(card):
        link_tag = card.find_parent('a', class_='card_link') or card.find('a', class_='card_link')
        program_link = ""
        if link_tag and link_tag.has_attr('href'):
            href = link_tag['href']
            if href.startswith('/'):
                href = "https://dpo.gubkin.ru" + href
            program_link = href
        title_elem = card.find('div', class_='card_title')
        title = title_elem.get_text(strip=True) if title_elem else ""
        direction_elem = card.find('div', class_='card_label_top')
        direction = direction_elem.get_text(strip=True) if direction_elem else ""
        formats = []
        formats_block = card.find('div', class_='card_title_add')
        if formats_block:
            for span in formats_block.find_all('span', class_='card_title_label'):
                formats.append(span.get_text(strip=True))
        formats_str = ', '.join(formats) if formats else ""
        info_block = card.find('div', class_='card_info')
        study_type = start_date = duration = cost = ""
        if info_block:
            lines = []
            current = []
            for child in info_block.children:
                if child.name == 'br':
                    lines.append(' '.join(current).strip())
                    current = []
                else:
                    text = child.get_text(strip=True) if hasattr(child, 'get_text') else str(child).strip()
                    if text:
                        current.append(text)
            if current:
                lines.append(' '.join(current).strip())
            for line in lines:
                if not line:
                    continue
                if 'повышение квалификации' in line or 'профессиональная переподготовка' in line:
                    study_type = line
                elif any(month in line for month in ['января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
                                                    'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря']) and 'с ' in line:
                    start_date = line
                elif any(key in line for key in ['недел', 'дн', 'месяц', 'час']):
                    duration = line
                elif 'руб' in line:
                    cost = line
        return {
            'Направление': direction, 'Название программы': title,
            'Ссылка на программу': program_link, 'Форматы': formats_str,
            'Вид обучения': study_type, 'Дата начала': start_date,
            'Длительность': duration, 'Стоимость': cost
        }

    print("=" * 60)
    print("ПАРСЕР ГУБКИНСКИЙ УНИВЕРСИТЕТ (dpo.gubkin.ru)")
    print("=" * 60)
    signal.signal(signal.SIGINT, signal_handler)
    driver = webdriver.Firefox()
    try:
        driver.get(URL)
        time.sleep(3)
        try:
            close_btn = WebDriverWait(driver, 5).until(
                EC.element_to_be_clickable((By.CSS_SELECTOR, "svg.svg-icon--IconClose"))
            )
            close_btn.click()
            time.sleep(1)
        except Exception:
            pass
        filter_buttons = driver.find_elements(By.CSS_SELECTOR, "div.myButton.bigbtn")
        print(f"Найдено фильтров: {len(filter_buttons)}")
        for idx, btn in enumerate(filter_buttons, 1):
            direction_name = btn.text.strip()
            if not direction_name:
                continue
            print(f"\n[{idx}/{len(filter_buttons)}] Обработка: {direction_name}")
            try:
                btn.click()
            except Exception as e:
                print(f"  Ошибка клика: {e}")
                continue
            try:
                WebDriverWait(driver, WAIT_TIMEOUT).until(
                    EC.presence_of_element_located((By.CSS_SELECTOR, "div.card"))
                )
            except Exception:
                print(f"  Карточки не появились за {WAIT_TIMEOUT} сек")
                time.sleep(DELAY_BETWEEN_FILTERS)
                continue
            time.sleep(1)
            if idx == len(filter_buttons):
                soup = BeautifulSoup(driver.page_source, 'html.parser')
                cards = soup.find_all('div', class_='card')
                print(f"  Найдено программ: {len(cards)}")
                for card in cards:
                    prog = parse_card(card)
                    last_direction_data.append(prog)
            time.sleep(DELAY_BETWEEN_FILTERS)
    finally:
        driver.quit()
    if last_direction_data:
        save_to_excel(last_direction_data, OUTPUT_FILENAME)
        print(f"Всего собрано записей: {len(last_direction_data)}")
    else:
        print("Данные не собраны")

# ========== ПАРСЕР №4: ОРЛОВСКИЙ УЧЕБНЫЙ КОМБИНАТ ==========

def run_parser_orlovskiy():
    """Парсер PDF с сайта uk57.ru."""
    OUTPUT_FILENAME = 'orlovskiy_data.xlsx'
    url = "https://uk57.ru/wp-content/uploads/2023/04/программа-обучения-1-1.pdf"
    print("=" * 60)
    print("ПАРСЕР ОРЛОВСКИЙ УЧЕБНЫЙ КОМБИНАТ (uk57.ru)")
    print("=" * 60)
    try:
        response = requests.get(url, verify=False, timeout=30)
        response.raise_for_status()
    except requests.exceptions.SSLError as e:
        print(f"Ошибка SSL: {e}")
        return
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        all_tables = []
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            for table in tables:
                if table and len(table) > 1:
                    headers = table[0]
                    data = table[1:]
                    data = [row for row in data if any(cell is not None and cell.strip() for cell in row)]
                    if data:
                        df = pd.DataFrame(data, columns=headers)
                        df['Page'] = page_num
                        all_tables.append(df)
        if all_tables:
            final_df = pd.concat(all_tables, ignore_index=True)
            final_df.to_excel(OUTPUT_FILENAME, index=False, engine='openpyxl')
            print(f"Файл {OUTPUT_FILENAME} успешно создан. Строк: {len(final_df)}")
        else:
            print("Таблицы в PDF не найдены")

# ========== ПАРСЕР №5: СТАНДАРТ ПЛЮС ==========

def run_parser_standart_plus():
    """Парсер с сайта standart82.ru (Selenium + авторизация)."""
    BASE_URL = "https://www.standart82.ru"
    LOGIN_URL = f"{BASE_URL}/login"
    SECTIONS = [
        {"name": "qualification-upgrade", "url": f"{BASE_URL}/qualification-upgrade"},
        {"name": "professional-retraining", "url": f"{BASE_URL}/professional-retraining"},
        {"name": "professional-education", "url": f"{BASE_URL}/professional-education"}
    ]
    USER_EMAIL = os.getenv('USER_EMAIL')
    USER_PASSWORD = os.getenv('USER_PASSWORD')
    DELAY = 1
    OUTPUT_FILENAME = 'standart_plus_data.xlsx'

    def setup_driver():
        firefox_options = FirefoxOptions()
        firefox_options.add_argument("--width=1920")
        firefox_options.add_argument("--height=1080")
        driver = webdriver.Firefox(options=firefox_options)
        driver.implicitly_wait(10)
        return driver

    def login(driver):
        print("Выполняется вход...")
        driver.get(LOGIN_URL)
        time.sleep(DELAY)
        try:
            email_field = WebDriverWait(driver, 15).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "input[type='email'], input[name='email']"))
            )
            password_field = driver.find_element(By.CSS_SELECTOR, "input[type='password'], input[name='password']")
            email_field.clear()
            email_field.send_keys(USER_EMAIL)
            time.sleep(DELAY)
            password_field.clear()
            password_field.send_keys(USER_PASSWORD)
            time.sleep(DELAY)
            login_button = driver.find_element(By.CSS_SELECTOR, "button[type='submit'], input[type='submit']")
            login_button.click()
            time.sleep(DELAY)
            WebDriverWait(driver, 15).until(EC.url_changes(LOGIN_URL))
            print("Вход выполнен успешно")
            time.sleep(DELAY)
            return True
        except Exception as e:
            print(f"Ошибка при авторизации: {e}")
            return False

    def collect_programs_from_section(driver, section):
        section_name = section['name']
        section_url = section['url']
        print(f"\nОбработка раздела: {section_name}")
        driver.get(section_url)
        time.sleep(DELAY)
        try:
            WebDriverWait(driver, 10).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "a[href*='/categories/']"))
            )
        except Exception:
            print("  Ссылки на категории не найдены")
            return []
        category_elements = driver.find_elements(By.CSS_SELECTOR, "a[href*='/categories/']")
        category_links = []
        for cat in category_elements:
            link = cat.get_attribute('href')
            if link and link not in category_links:
                category_links.append(link)
        print(f"  Найдено категорий: {len(category_links)}")
        if not category_links:
            return []
        all_rows = []
        for idx, cat_link in enumerate(category_links, 1):
            print(f"  Категория {idx}/{len(category_links)}: {cat_link}")
            driver.get(cat_link)
            time.sleep(DELAY)
            soup = BeautifulSoup(driver.page_source, 'lxml')
            program_cards = []
            title_tags = soup.find_all('h3', class_='line-clamp-2')
            for title_tag in title_tags:
                card = title_tag.find_parent('div', class_='relative')
                if not card:
                    card = title_tag.find_parent('div', recursive=False)
                if card:
                    program_cards.append(card)
            if not program_cards:
                print("    Не найдено карточек программ")
                continue
            for card in program_cards:
                title_tag = card.find('h3', class_='line-clamp-2')
                title = title_tag.get_text(strip=True) if title_tag else ''
                price = 'Цена не указана'
                price_part = card.find(string=re.compile(r'[\d\s]+₽'))
                if price_part:
                    price = price_part.strip()
                else:
                    price_div = card.find('div', class_='flex', string=re.compile(r'₽'))
                    if price_div:
                        price = price_div.get_text(strip=True)
                hours_span = card.find('span', string=re.compile(r'\d+ч'))
                hours = hours_span.get_text(strip=True) if hours_span else ''
                link_tag = card.find_parent('a', href=re.compile(r'/programs/'))
                if not link_tag:
                    link_tag = card.find('a', href=re.compile(r'/programs/'))
                program_url = urljoin(BASE_URL, link_tag['href']) if link_tag else None
                all_rows.append({
                    'Раздел': section_name,
                    'Категория': cat_link.split('/')[-1],
                    'Название программы': title,
                    'Цена': price,
                    'Часы': hours,
                    'Ссылка': program_url
                })
            print(f"    Найдено программ: {len(program_cards)}")
            time.sleep(DELAY)
        return all_rows

    print("=" * 60)
    print("ПАРСЕР СТАНДАРТ ПЛЮС (standart82.ru)")
    print("=" * 60)
    if not USER_EMAIL or not USER_PASSWORD:
        print("Ошибка: не заданы USER_EMAIL / USER_PASSWORD в .env")
        return
    driver = setup_driver()
    try:
        if not login(driver):
            return
        all_data = []
        for section in SECTIONS:
            records = collect_programs_from_section(driver, section)
            all_data.extend(records)
            time.sleep(DELAY)
        if all_data:
            df = pd.DataFrame(all_data)
            df.to_excel(OUTPUT_FILENAME, index=False, engine='openpyxl')
            print(f"Данные сохранены в {OUTPUT_FILENAME}")
            print(f"Всего собрано программ: {len(all_data)}")
        else:
            print("Нет данных для сохранения")
    finally:
        driver.quit()
        print("Браузер закрыт")

# ========== ПАРСЕР №6: АКАДЕМИЯ (АНО УЦ ДПО) ==========

def run_parser_academia():
    """Парсер PDF с сайта anodpo.ru."""
    OUTPUT_FILENAME = 'academia_data.xlsx'
    url = "https://anodpo.ru/price.pdf"
    print("=" * 60)
    print("ПАРСЕР АКАДЕМИЯ (anodpo.ru)")
    print("=" * 60)
    print("Загружаем PDF...")
    resp = requests.get(url, verify=False)
    resp.raise_for_status()
    with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as tmp:
        tmp.write(resp.content)
        tmp_path = tmp.name
    try:
        print("Извлекаем таблицы (flavor='lattice')...")
        tables = camelot.read_pdf(tmp_path, pages='4-end', flavor='lattice')
        print(f"Найдено таблиц: {len(tables)}")
        all_dfs = []
        for table in tables:
            df = table.df.copy()
            df.insert(0, 'Страница', table.page)
            all_dfs.append(df)
        final_df = pd.concat(all_dfs, ignore_index=True, sort=False)
        final_df.to_excel(OUTPUT_FILENAME, index=False)
        print(f"Сохранено {len(final_df)} строк в {OUTPUT_FILENAME}")
    except Exception as e:
        print(f"Ошибка: {e}")
        import traceback
        traceback.print_exc()
    finally:
        try:
            os.unlink(tmp_path)
        except PermissionError:
            pass

# ========== ПАРСЕР №7: МЭИ ==========

def run_parser_mpei():
    """Парсер с сайта mpei.ru."""
    base_url = 'https://mpei.ru'
    url = base_url + '/Education/AdditionalEducation/Pages/programs.aspx'
    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
    OUTPUT_FILENAME = 'mpei_data.xlsx'
    print("=" * 60)
    print("ПАРСЕР МЭИ (mpei.ru)")
    print("=" * 60)
    try:
        response = requests.get(url, headers=headers, verify=False, timeout=30)
        response.raise_for_status()
        response.encoding = 'utf-8'
    except requests.exceptions.RequestException as e:
        print(f"Ошибка при загрузке страницы: {e}")
        return
    soup = BeautifulSoup(response.text, 'html.parser')
    direction_headers = soup.find_all('div', class_='groupheader')
    print(f"Найдено направлений: {len(direction_headers)}")
    all_programs = []
    for header in direction_headers:
        direction_name = header.get_text(strip=True)
        parent_li = header.find_parent('li', class_='dfwp-item')
        if not parent_li:
            continue
        tables = parent_li.find_all('table', class_='DopProgList')
        for table in tables:
            rows = table.find_all('tr')
            for row in rows:
                if row.find('th', class_='DopProgListH'):
                    continue
                cells = row.find_all('td')
                if len(cells) < 5:
                    continue
                code = cells[0].get_text(strip=True)
                if not code:
                    continue
                program_tag = cells[1].find('a')
                if program_tag:
                    program_name = program_tag.get_text(strip=True)
                    program_link = urljoin(base_url, program_tag.get('href', ''))
                else:
                    program_name = cells[1].get_text(strip=True)
                    program_link = ''
                form = cells[2].get_text(strip=True)
                direction_field = cells[3].get_text(strip=True)
                department = cells[4].get_text(strip=True)
                all_programs.append({
                    'Код направления': code, 'Профиль / Программа': program_name,
                    'Форма обучения': form, 'Направление': direction_name,
                    'Кафедра': department, 'Ссылка на программу': program_link
                })
    print(f"Всего собрано записей: {len(all_programs)}")
    if not all_programs:
        print("Данные не найдены")
        return
    df = pd.DataFrame(all_programs)
    df.to_excel(OUTPUT_FILENAME, index=False, engine='openpyxl')
    print(f"Данные сохранены в {OUTPUT_FILENAME}")

# ========== ПАРСЕР №8: ЭНЕРГЕТИК (АНО ДПО) ==========

def run_parser_energetik_dpo():
    """Парсер с сайта dpo-energo.ru (Selenium)."""
    URL = "https://dpo-energo.ru/education/"
    OUTPUT_FILENAME = 'energetik_dpo_data.xlsx'
    DELAY_BETWEEN_SECTIONS = 2
    WAIT_TIMEOUT = 10
    last_section_data = []

    def signal_handler(sig, frame):
        print("\nПрерывание. Сохраняем собранные данные...")
        if last_section_data:
            save_to_excel(last_section_data, OUTPUT_FILENAME)
        sys.exit(0)

    def save_to_excel(data, filename):
        if not data:
            return
        if os.path.exists(filename):
            try:
                existing_df = pd.read_excel(filename)
                existing_data = existing_df.to_dict('records')
            except Exception as e:
                print(f"Не удалось загрузить существующий файл: {e}")
                existing_data = []
        else:
            existing_data = []
        combined = existing_data + data
        df = pd.DataFrame(combined)
        df = df.drop_duplicates(subset=['Название курса', 'Количество часов'], keep='first')
        df = df.sort_values('Название курса')
        df.to_excel(filename, index=False)
        print(f"Сохранено {len(df)} записей в {filename}")

    def parse_course_row(row, section_name=""):
        title_div = row.find('div', class_='font-bold')
        if not title_div:
            return None
        course_name = title_div.get_text(strip=True)
        details_div = row.find('div', class_=re.compile(r'text-\[10px\]'))
        if not details_div:
            return None
        details_text = details_div.get_text(strip=True)
        hours = ""
        form = ""
        hours_match = re.search(r'(\d+)\s*ч\.', details_text)
        if hours_match:
            hours = hours_match.group(1) + ' ч.'
        form_match = re.search(r'Форма\s*-\s*([а-яА-ЯёЁ\s/]+)', details_text)
        if form_match:
            form = form_match.group(1).strip()
        return {
            'Название курса': course_name, 'Количество часов': hours,
            'Форма обучения': form, 'Направление / секция': section_name
        }

    def get_section_name_from_button(driver, button):
        try:
            section = driver.execute_script("return arguments[0].closest('section')", button)
            if section:
                h2 = section.find_element(By.TAG_NAME, "h2")
                return h2.text.strip()
        except Exception:
            pass
        try:
            span = button.find_element(By.CSS_SELECTOR, "span.text-xs.font-black")
            return span.text.strip()
        except Exception:
            return "Неизвестная секция"

    print("=" * 60)
    print("ПАРСЕР ЭНЕРГЕТИК (dpo-energo.ru)")
    print("=" * 60)
    signal.signal(signal.SIGINT, signal_handler)
    options = webdriver.FirefoxOptions()
    driver = webdriver.Firefox(options=options)
    try:
        driver.get(URL)
        time.sleep(3)
        try:
            close_btn = WebDriverWait(driver, 5).until(
                EC.element_to_be_clickable((By.CSS_SELECTOR, "button[aria-label='Закрыть'], .modal-close"))
            )
            close_btn.click()
            time.sleep(1)
        except Exception:
            pass
        buttons = driver.find_elements(By.CSS_SELECTOR, "button[onclick*='toggleCourses']")
        total_buttons = len(buttons)
        print(f"Найдено секций: {total_buttons}")
        buttons_info = []
        for btn in buttons:
            onclick = btn.get_attribute('onclick')
            match = re.search(r"toggleCourses\('([^']+)'\)", onclick)
            if match:
                block_id = match.group(1)
                section_name = get_section_name_from_button(driver, btn)
                buttons_info.append({'block_id': block_id, 'section_name': section_name})
        for idx, info in enumerate(buttons_info, 1):
            block_id = info['block_id']
            section_name = info['section_name']
            print(f"\n[{idx}/{total_buttons}] Обработка: {section_name} (ID: {block_id})")
            try:
                btn = driver.find_element(By.CSS_SELECTOR, f"button[onclick*='{block_id}']")
            except Exception as e:
                print(f"  Не удалось найти кнопку: {e}")
                continue
            driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", btn)
            time.sleep(0.5)
            try:
                driver.execute_script("arguments[0].click();", btn)
            except Exception as e:
                print(f"  Ошибка клика: {e}")
                continue
            try:
                WebDriverWait(driver, WAIT_TIMEOUT).until(
                    EC.visibility_of_element_located((By.ID, block_id))
                )
            except Exception as e:
                print(f"  Блок {block_id} не появился: {e}")
                time.sleep(DELAY_BETWEEN_SECTIONS)
                continue
            time.sleep(1)
            block_element = driver.find_element(By.ID, block_id)
            block_html = block_element.get_attribute('innerHTML')
            soup_block = BeautifulSoup(block_html, 'html.parser')
            table = soup_block.find('table', class_=re.compile('border-collapse'))
            if not table:
                print("  Таблица не найдена")
                time.sleep(DELAY_BETWEEN_SECTIONS)
                continue
            rows = table.find_all('tr')
            courses_count = 0
            for row in rows:
                course_data = parse_course_row(row, section_name)
                if course_data:
                    last_section_data.append(course_data)
                    courses_count += 1
            print(f"  Добавлено курсов: {courses_count}")
            time.sleep(DELAY_BETWEEN_SECTIONS)
    finally:
        driver.quit()
    if last_section_data:
        save_to_excel(last_section_data, OUTPUT_FILENAME)
        print(f"Всего собрано записей: {len(last_section_data)}")
    else:
        print("Данные не собраны")

# ========== ПАРСЕР №9: ЕИПК ==========

def run_parser_eipk():
    """Парсер с сайта eipk.ru."""
    url = "https://eipk.ru/about/learn/"
    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
    OUTPUT_FILENAME = 'eipk_data.xlsx'
    print("=" * 60)
    print("ПАРСЕР ЕИПК (eipk.ru)")
    print("=" * 60)
    try:
        response = requests.get(url, headers=headers, verify=False)
        response.raise_for_status()
        response.encoding = 'utf-8'
    except requests.exceptions.RequestException as e:
        print(f"Ошибка при загрузке страницы: {e}")
        return
    soup = BeautifulSoup(response.text, 'html.parser')
    sections = soup.find_all('div', class_='sectionW')
    if not sections:
        print("Не удалось найти секции на странице")
        return
    all_data = []
    for section in sections:
        section_header = section.find('div', class_='sectionH')
        section_name = section_header.get_text(strip=True) if section_header else "Название секции не найдено"
        table = section.find('table', class_='table-style-1')
        if not table:
            continue
        tbody = table.find('tbody')
        if not tbody:
            tbody = table
        rows = tbody.find_all('tr')
        for row in rows:
            cols = row.find_all('td')
            if len(cols) >= 4:
                program_number = cols[0].get_text(strip=True)
                link_tag = cols[1].find('a')
                program_name = link_tag.get_text(strip=True) if link_tag else cols[1].get_text(strip=True)
                program_url = urljoin(url, link_tag['href']) if link_tag and link_tag.has_attr('href') else ""
                education_type = cols[2].get_text(strip=True)
                hours = cols[3].get_text(strip=True)
                all_data.append({
                    'Секция': section_name, '№ п/п': program_number,
                    'Наименование программы': program_name, 'Ссылка': program_url,
                    'Вид образования': education_type, 'Количество часов': hours
                })
    if all_data:
        df = pd.DataFrame(all_data)
        df.to_excel(OUTPUT_FILENAME, index=False, engine='openpyxl')
        print(f"Данные сохранены в {OUTPUT_FILENAME}. Записей: {len(df)}")
    else:
        print("Не удалось найти данные для парсинга")

# ========== ПАРСЕР №10: ЕЦОИПС ==========

def run_parser_ecoips():
    """Парсер с сайта ecoips.ru."""
    BASE_URL = "https://ecoips.ru"
    MAIN_URL = "https://ecoips.ru/obuchenie-i-uslugi/"
    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
    OUTPUT_FILENAME = 'ecoips_data.xlsx'

    def get_soup(url):
        try:
            response = requests.get(url, headers=headers, timeout=10, verify=False)
            response.raise_for_status()
            response.encoding = 'utf-8'
            return BeautifulSoup(response.text, 'html.parser')
        except requests.exceptions.RequestException as e:
            print(f"Ошибка загрузки {url}: {e}")
            return None

    def get_program_hours(program_url):
        soup = get_soup(program_url)
        if not soup:
            return None
        for text_elem in soup.find_all(text=True):
            if "Продолжительность курса:" in text_elem:
                parent = text_elem.find_parent()
                if parent:
                    text = parent.get_text()
                    match = re.search(r'Продолжительность курса:\s*(\d+)', text)
                    if match:
                        return match.group(1)
                break
        return None

    print("=" * 60)
    print("ПАРСЕР ЕЦОИПС (ecoips.ru)")
    print("=" * 60)
    print("Шаг 1: Парсинг главной страницы...")
    main_soup = get_soup(MAIN_URL)
    if not main_soup:
        print("Не удалось загрузить главную страницу")
        return
    sections = main_soup.select('.box a.item')
    print(f"Найдено разделов: {len(sections)}")
    all_courses = []
    for section in sections:
        section_name_elem = section.select_one('.name')
        if not section_name_elem:
            continue
        section_name = section_name_elem.get_text(strip=True)
        section_url = urljoin(BASE_URL, section.get('href'))
        print(f"\nОбработка раздела: {section_name} ({section_url})")
        section_soup = get_soup(section_url)
        if not section_soup:
            continue
        programs_container = section_soup.find('div', id='content')
        if not programs_container:
            continue
        program_links = programs_container.select('a.item')
        print(f"Найдено программ в разделе: {len(program_links)}")
        for program in program_links:
            program_name_elem = program.select_one('.zag')
            hours_elem = program.select_one('.date')
            if not program_name_elem:
                continue
            program_name = program_name_elem.get_text(strip=True)
            hours_on_page = hours_elem.get_text(strip=True) if hours_elem else None
            program_url = urljoin(BASE_URL, program.get('href'))
            program_hours = hours_on_page
            if not program_hours and program_url:
                program_hours = get_program_hours(program_url)
            all_courses.append({
                'Раздел': section_name, 'Программа': program_name,
                'Часы': program_hours, 'Ссылка': program_url
            })
    print(f"\nСобрано записей: {len(all_courses)}")
    df = pd.DataFrame(all_courses)
    df.to_excel(OUTPUT_FILENAME, index=False)
    print(f"Данные сохранены в {OUTPUT_FILENAME}")

# ========== ПАРСЕР №11: НЕФТЕГАЗОВЫЙ ОБРАЗОВАТЕЛЬНЫЙ ЦЕНТР ==========

def run_parser_neftegaz():
    """Парсер PDF с сайта oilgasec.ru."""
    OUTPUT_FILENAME = 'neftegaz_data.xlsx'
    url = "https://oilgasec.ru/wp-content/uploads/Прайс-на-образовательные-услуги-2026_НОЦ-1.pdf"
    print("=" * 60)
    print("ПАРСЕР НЕФТЕГАЗОВЫЙ ОБРАЗОВАТЕЛЬНЫЙ ЦЕНТР (oilgasec.ru)")
    print("=" * 60)
    print("Загружаем PDF...")
    resp = requests.get(url, verify=False)
    resp.raise_for_status()
    with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as tmp:
        tmp.write(resp.content)
        tmp_path = tmp.name
    try:
        print("Извлекаем таблицы (flavor='lattice')...")
        tables = camelot.read_pdf(tmp_path, pages='2-end', flavor='lattice')
        print(f"Найдено таблиц: {len(tables)}")
        all_dfs = []
        for table in tables:
            df = table.df.copy()
            df.insert(0, 'Страница', table.page)
            df.columns = range(df.shape[1])
            all_dfs.append(df)
        final_df = pd.concat(all_dfs, ignore_index=True, sort=False)
        headers_row = final_df.iloc[0, 1:]
        new_columns = ['Страница'] + headers_row.tolist()
        final_df.columns = new_columns
        final_df = final_df[1:].reset_index(drop=True)
        if '№ п/п' in final_df.columns:
            mask = (final_df['№ п/п'] != '№ п/п') & (final_df['№ п/п'].notna()) & (final_df['№ п/п'] != '')
            final_df = final_df[mask].reset_index(drop=True)
        final_df.to_excel(OUTPUT_FILENAME, index=False)
        print(f"Сохранено {len(final_df)} строк в {OUTPUT_FILENAME}")
    except Exception as e:
        print(f"Ошибка: {e}")
        import traceback
        traceback.print_exc()
    finally:
        try:
            os.unlink(tmp_path)
        except PermissionError:
            pass

# ========== ПАРСЕР №12: ТПУ ==========

def run_parser_tpu():
    """Парсер с сайта dpo.tpu.ru (Selenium)."""
    URL = "https://dpo.tpu.ru/courses/"
    OUTPUT_FILENAME = 'tpu_data.xlsx'
    DELAY_BETWEEN_TABS = 2
    WAIT_TIMEOUT = 10
    SCROLL_PAUSE_TIME = 1
    MAX_SCROLL_ATTEMPTS = 30
    all_programs = []

    def signal_handler(sig, frame):
        print("\nПрерывание. Сохраняем собранные данные...")
        if all_programs:
            save_to_excel(all_programs, OUTPUT_FILENAME)
        sys.exit(0)

    def save_to_excel(data, filename):
        if not data:
            return
        if os.path.exists(filename):
            try:
                existing_df = pd.read_excel(filename)
                existing_data = existing_df.to_dict('records')
            except Exception as e:
                print(f"Не удалось загрузить существующий файл: {e}")
                existing_data = []
        else:
            existing_data = []
        combined = existing_data + data
        df = pd.DataFrame(combined)
        df = df.drop_duplicates(subset=['Раздел', 'Название программы', 'Часы', 'Стоимость'], keep='first')
        cols_order = ['Раздел', 'Название программы', 'Ссылка', 'Часы', 'Стоимость', 'Комментарий']
        df = df[cols_order]
        df = df.sort_values('Раздел')
        df.to_excel(filename, index=False)
        print(f"Сохранено {len(df)} записей в {filename}")

    def parse_card(card, section_name):
        hours_elem = card.find('div', class_='direction__training')
        hours = ""
        if hours_elem:
            short_title = hours_elem.find('span', class_='short-title')
            if short_title:
                hours = short_title.get_text(strip=True)
        price = ""
        if hours_elem:
            title_elem = hours_elem.find('span', class_='title')
            if title_elem:
                price = title_elem.get_text(strip=True)
        prog_elem = card.find('div', class_='direction__educational-program')
        comment = ""
        if prog_elem:
            comment_elem = prog_elem.find('span', class_='short-title')
            if comment_elem:
                comment = comment_elem.get_text(strip=True)
        program_name = ""
        if prog_elem:
            name_elem = prog_elem.find('h4')
            if name_elem:
                a_tag = name_elem.find('a')
                if a_tag:
                    program_name = a_tag.get_text(strip=True)
        link = ""
        more_link = card.find('a', class_='more')
        if more_link and more_link.has_attr('href'):
            link = more_link['href']
        return {
            'Раздел': section_name, 'Название программы': program_name,
            'Ссылка': link, 'Часы': hours, 'Стоимость': price, 'Комментарий': comment
        }

    def scroll_to_bottom(driver, pause=SCROLL_PAUSE_TIME, max_attempts=MAX_SCROLL_ATTEMPTS):
        last_height = driver.execute_script("return document.body.scrollHeight")
        for _ in range(max_attempts):
            driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
            time.sleep(pause)
            new_height = driver.execute_script("return document.body.scrollHeight")
            if new_height == last_height:
                driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                time.sleep(pause)
                new_height = driver.execute_script("return document.body.scrollHeight")
                if new_height == last_height:
                    break
            last_height = new_height

    print("=" * 60)
    print("ПАРСЕР ТПУ (dpo.tpu.ru)")
    print("=" * 60)
    signal.signal(signal.SIGINT, signal_handler)
    driver = webdriver.Firefox()
    try:
        driver.get(URL)
        time.sleep(3)
        try:
            close_btn = WebDriverWait(driver, 5).until(
                EC.element_to_be_clickable((By.CSS_SELECTOR, "button.modal__close"))
            )
            close_btn.click()
            time.sleep(1)
        except Exception:
            pass
        tab_buttons = driver.find_elements(By.CSS_SELECTOR, "button.tabs-navigation__btn")
        print(f"Найдено вкладок: {len(tab_buttons)}")
        for idx, btn in enumerate(tab_buttons):
            section_name = btn.text.strip()
            if not section_name:
                continue
            print(f"\n[{idx + 1}/{len(tab_buttons)}] Обработка: {section_name}")
            try:
                driver.execute_script("arguments[0].click();", btn)
            except Exception as e:
                print(f"  Ошибка клика: {e}")
                continue
            time.sleep(DELAY_BETWEEN_TABS)
            scroll_to_bottom(driver)
            time.sleep(SCROLL_PAUSE_TIME)
            tab_content_containers = driver.find_elements(By.CSS_SELECTOR, "div.js-tab")
            if idx >= len(tab_content_containers):
                continue
            content_el = tab_content_containers[idx]
            html = content_el.get_attribute('outerHTML')
            soup = BeautifulSoup(html, 'html.parser')
            cards = soup.find_all('div', class_='direction')
            print(f"  Найдено программ: {len(cards)}")
            for card in cards:
                data = parse_card(card, section_name)
                if data['Название программы']:
                    all_programs.append(data)
            time.sleep(DELAY_BETWEEN_TABS)
    finally:
        driver.quit()
    if all_programs:
        save_to_excel(all_programs, OUTPUT_FILENAME)
        print(f"Всего собрано записей: {len(all_programs)}")
    else:
        print("Данные не собраны")

# ========== ПАРСЕР №13: АЗ ==========

def run_parser_az():
    """Парсер с сайта az30.ru."""
    from urllib3.exceptions import InsecureRequestWarning
    from requests.adapters import HTTPAdapter
    from urllib3.util.retry import Retry
    from requests.exceptions import ReadTimeout
    requests.packages.urllib3.disable_warnings(InsecureRequestWarning)
    HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
    BASE_URL = 'https://az30.ru/shop/'
    OUTPUT_FILENAME = 'az_data.xlsx'
    session = requests.Session()
    session.verify = False
    session.headers.update(HEADERS)
    retry_strategy = Retry(total=3, backoff_factor=1, status_forcelist=[500, 502, 503, 504], allowed_methods=["HEAD", "GET", "OPTIONS"])
    adapter = HTTPAdapter(max_retries=retry_strategy)
    session.mount("http://", adapter)
    session.mount("https://", adapter)

    def get_soup(url):
        try:
            response = session.get(url, timeout=30)
            response.raise_for_status()
            response.encoding = 'utf-8'
            return BeautifulSoup(response.text, 'html.parser')
        except ReadTimeout:
            print(f'Таймаут при загрузке {url}')
            return None
        except Exception as e:
            print(f'Ошибка при загрузке {url}: {e}')
            return None

    def clean_category_name(name):
        return re.sub(r'\s*\(\d+\)\s*$', '', name).strip()

    def extract_hours(title):
        match = re.search(r'\(([^)]+)\)', title)
        if match:
            hours = match.group(1).strip()
            clean_title = re.sub(r'\s*\([^)]+\)\s*$', '', title).strip()
            return clean_title, hours
        return title, ''

    def parse_categories(soup):
        categories = []
        if not soup:
            return categories
        product_cats = soup.find_all('li', class_='product-category')
        for cat in product_cats:
            a_tag = cat.find('a')
            if not a_tag:
                continue
            href = a_tag.get('href')
            if not href:
                continue
            h2 = cat.find('h2', class_='woocommerce-loop-category__title')
            if h2:
                raw_name = h2.get_text(strip=True)
                name = clean_category_name(raw_name)
            else:
                name = 'Без названия'
            categories.append({'name': name, 'url': href})
        return categories

    def parse_prices_from_table(product_soup):
        common_prices = {'индивидуально': None, 'очный курс': None, 'очно-дистанционный курс': None, 'дистанционный курс': None}
        detailed_prices = {
            'подготовка_индивидуально': None, 'подготовка_очный': None,
            'подготовка_очно-дистанционный': None, 'подготовка_дистанционный': None,
            'повышение_очный': None, 'повышение_очно-дистанционный': None,
            'повышение_дистанционный': None, 'повышение_индивидуально': None
        }
        table = product_soup.find('table', class_='az30-variations-list')
        if not table:
            return common_prices, detailed_prices
        rows = table.find_all('tr')
        for row in rows:
            cells = row.find_all('td')
            if len(cells) < 2:
                continue
            type_text = cells[0].get_text(strip=True).lower()
            price_text = cells[1].get_text(strip=True)
            price_match = re.search(r'([\d\s,]+)', price_text.replace('₽', ''))
            price = None
            if price_match:
                price_str = price_match.group(1).replace(' ', '').replace(',', '.')
                try:
                    price = float(price_str)
                except Exception:
                    price = None
            if common_prices['индивидуально'] is None and 'индивидуальн' in type_text:
                common_prices['индивидуально'] = price
            if common_prices['очный курс'] is None and 'очный курс' in type_text:
                common_prices['очный курс'] = price
            if common_prices['очно-дистанционный курс'] is None and 'очно-дистанцион' in type_text:
                common_prices['очно-дистанционный курс'] = price
            if common_prices['дистанционный курс'] is None and 'дистанционн' in type_text:
                common_prices['дистанционный курс'] = price
            if 'подготовка/переподготовка индивидуально' in type_text:
                detailed_prices['подготовка_индивидуально'] = price
            elif 'подготовка/переподготовка очно' in type_text:
                detailed_prices['подготовка_очный'] = price
            elif 'подготовка/переподготовка очно-дистанционно' in type_text:
                detailed_prices['подготовка_очно-дистанционный'] = price
            elif 'подготовка/переподготовка дистанционно' in type_text:
                detailed_prices['подготовка_дистанционный'] = price
            elif 'повышение квалификации очно' in type_text:
                detailed_prices['повышение_очный'] = price
            elif 'повышение квалификации очно-дистанционно' in type_text:
                detailed_prices['повышение_очно-дистанционный'] = price
            elif 'повышение квалификации дистанционно' in type_text:
                detailed_prices['повышение_дистанционный'] = price
            elif 'повышение квалификации индивидуально' in type_text:
                detailed_prices['повышение_индивидуально'] = price
        return common_prices, detailed_prices

    def parse_product_card(product_li, category_name):
        a_tag = product_li.find('a', class_='woocommerce-LoopProduct-link')
        if not a_tag:
            return None
        product_url = a_tag.get('href')
        h2 = a_tag.find('h2', class_='woocommerce-loop-product__title')
        if not h2:
            return None
        full_title = h2.get_text(strip=True)
        clean_title, hours = extract_hours(full_title)
        common_prices, detailed_prices = parse_prices_from_table(product_li)
        price_span = product_li.find('span', class_='price')
        min_price = max_price = None
        if price_span:
            span_text = price_span.get_text(strip=True)
            parts = re.split(r'[–-]', span_text)
            if len(parts) >= 2:
                try:
                    min_price = float(re.search(r'([\d\s,]+\.?\d*)', parts[0].replace('₽', '')).group(1).replace(' ', '').replace(',', '.'))
                except Exception:
                    pass
                try:
                    max_price = float(re.search(r'([\d\s,]+\.?\d*)', parts[1].replace('₽', '')).group(1).replace(' ', '').replace(',', '.'))
                except Exception:
                    pass
            else:
                try:
                    price_val = float(re.search(r'([\d\s,]+\.?\d*)', span_text.replace('₽', '')).group(1).replace(' ', '').replace(',', '.'))
                    min_price = max_price = price_val
                except Exception:
                    pass
        if min_price is None:
            all_vals = list(common_prices.values()) + list(detailed_prices.values())
            valid = [v for v in all_vals if v is not None]
            if valid:
                min_price = min(valid)
                max_price = max(valid)
        return {
            'Раздел': category_name, 'Наименование программы': clean_title,
            'количество часов': hours, 'сумма от': min_price, 'сумма до': max_price,
            'индивидуально': common_prices['индивидуально'],
            'очный курс': common_prices['очный курс'],
            'очно-дистанционный курс': common_prices['очно-дистанционный курс'],
            'дистанционный курс': common_prices['дистанционный курс'],
            'ссылка на программу': product_url,
            'Подготовка/переподготовка индивидуально': detailed_prices['подготовка_индивидуально'],
            'Подготовка/переподготовка очно': detailed_prices['подготовка_очный'],
            'Подготовка/переподготовка очно-дистанционно': detailed_prices['подготовка_очно-дистанционный'],
            'Подготовка/переподготовка дистанционно': detailed_prices['подготовка_дистанционный'],
            'Повышение квалификации очно': detailed_prices['повышение_очный'],
            'Повышение квалификации очно-дистанционно': detailed_prices['повышение_очно-дистанционный'],
            'Повышение квалификации дистанционно': detailed_prices['повышение_дистанционный'],
            'Повышение квалификации индивидуально': detailed_prices['повышение_индивидуально']
        }

    def parse_category_products(category_url, category_name):
        all_products = []
        page = 1
        while True:
            url = category_url if page == 1 else f"{category_url.rstrip('/')}/page/{page}/"
            print(f"  Загрузка страницы {page}: {url}")
            soup = get_soup(url)
            if not soup:
                break
            products = soup.find_all('li', class_='product')
            if not products:
                break
            print(f"  Найдено товаров: {len(products)}")
            for prod in products:
                data = parse_product_card(prod, category_name)
                if data:
                    all_products.append(data)
            pagination = soup.find('ul', class_='page-numbers')
            if pagination:
                next_link = pagination.find('a', class_='next')
                if not next_link:
                    break
            else:
                break
            page += 1
            time.sleep(1)
        return all_products

    print("=" * 60)
    print("ПАРСЕР АЗ (az30.ru)")
    print("=" * 60)
    main_soup = get_soup(BASE_URL)
    if not main_soup:
        return
    categories = parse_categories(main_soup)
    print(f"Найдено категорий: {len(categories)}")
    all_data = []
    for idx, cat in enumerate(categories, 1):
        print(f"\n[{idx}/{len(categories)}] Обработка категории: {cat['name']}")
        products = parse_category_products(cat['url'], cat['name'])
        all_data.extend(products)
        time.sleep(1)
    df = pd.DataFrame(all_data)
    columns_order = [
        'Раздел', 'Наименование программы', 'количество часов', 'сумма от', 'сумма до',
        'индивидуально', 'очный курс', 'очно-дистанционный курс', 'дистанционный курс',
        'Подготовка/переподготовка индивидуально', 'Подготовка/переподготовка очно',
        'Подготовка/переподготовка очно-дистанционно', 'Подготовка/переподготовка дистанционно',
        'Повышение квалификации очно', 'Повышение квалификации очно-дистанционно',
        'Повышение квалификации дистанционно', 'Повышение квалификации индивидуально',
        'ссылка на программу'
    ]
    df = df[[col for col in columns_order if col in df.columns]]
    df.to_excel(OUTPUT_FILENAME, index=False, engine='openpyxl')
    print(f"Сохранено {len(df)} записей в {OUTPUT_FILENAME}")

# ========== ПАРСЕР №14: ЦППК ==========

def run_parser_cppk():
    """Парсер с сайта cpkrus.ru."""
    from openpyxl import Workbook
    BASE_URL = 'https://cpkrus.ru'
    HEADERS = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8',
        'Accept-Language': 'ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7',
    }
    OUTPUT_FILENAME = 'cppk_data.xlsx'

    def get_soup(url):
        response = requests.get(url, headers=HEADERS, verify=False, timeout=15)
        response.encoding = 'utf-8'
        return BeautifulSoup(response.text, 'html.parser')

    def get_categories():
        soup = get_soup('https://cpkrus.ru/tyumen/kursy-professionalnoj-perepodgotovki/')
        categories = []
        tiles = soup.select('div.directions-module__item.direction-tile')
        print(f"Найдено блоков категорий: {len(tiles)}")
        for tile in tiles:
            detail_button = tile.select_one('a.custom-button--gray-o')
            if not detail_button or not detail_button.get('href'):
                continue
            name_tag = tile.select_one('div.direction-tile__name a')
            name = name_tag.get_text(strip=True) if name_tag else 'Без названия'
            link = detail_button.get('href')
            full_link = urljoin(BASE_URL, link)
            categories.append({'name': name, 'url': full_link})
            print(f"  Найдена категория: {name}")
        return categories

    def extract_char_from_course_tile(course, label):
        label_elem = course.find('div', class_='course-tile__characteristics-label', string=lambda t: t and label in t)
        if label_elem:
            value_elem = label_elem.find_next_sibling('div', class_='course-tile__characteristics-value')
            if value_elem:
                return value_elem.get_text(strip=True)
        return ''

    def parse_courses_from_tiles(soup):
        courses = []
        for course in soup.select('div.course-tile'):
            title_tag = course.select_one('div.course-tile__title a')
            if not title_tag:
                continue
            title = title_tag.get_text(strip=True)
            link = title_tag.get('href')
            full_link = urljoin(BASE_URL, link)
            volume = extract_char_from_course_tile(course, 'Объем программы')
            duration = extract_char_from_course_tile(course, 'Срок обучения')
            document = extract_char_from_course_tile(course, 'Получаемый документ')
            education = extract_char_from_course_tile(course, 'Требуемое образование')
            price_full = ''
            price_block = course.select_one('.course-tile__prices--body:first-child .course-tile__prices-current')
            if price_block:
                price_full = price_block.get_text(strip=True)
            courses.append({
                'title': title, 'volume': volume, 'duration': duration,
                'document': document, 'education': education,
                'price_full': price_full, 'url': full_link
            })
        return courses

    def parse_single_program_page(soup, category_url):
        courses = []
        title_tag = soup.select_one('h1')
        title = title_tag.get_text(strip=True) if title_tag else 'Без названия'
        full_link = category_url
        volume_tabs = soup.select('.program-volume__tab')
        volume = ', '.join([tab.get_text(strip=True) for tab in volume_tabs]) if volume_tabs else ''
        duration = ''
        document = ''
        doc_panel = soup.select_one('.documents-section__panel--active .documents-section__description')
        if doc_panel and 'диплом' in doc_panel.get_text().lower():
            document = 'Диплом о профессиональной переподготовке'
        else:
            document = 'Диплом'
        education = ''
        edu_text = soup.find(string=lambda t: t and 'Требуемое образование' in t)
        if edu_text:
            parent = edu_text.find_parent()
            if parent:
                value = parent.find_next_sibling()
                if value:
                    education = value.get_text(strip=True)
        if not education:
            education = 'Высшее или среднее профессиональное'
        price_full = ''
        active_price = soup.select_one('.course-price__content.active .course-price__price')
        if active_price:
            price_full = active_price.get_text(strip=True)
        courses.append({
            'title': title, 'volume': volume, 'duration': duration,
            'document': document, 'education': education,
            'price_full': price_full, 'url': full_link
        })
        return courses

    def get_courses_from_category(category_url):
        soup = get_soup(category_url)
        if soup.select('div.course-tile'):
            return parse_courses_from_tiles(soup)
        elif soup.select_one('.courseinfo-wrap'):
            return parse_single_program_page(soup, category_url)
        else:
            print(f"  Не удалось найти программы на {category_url}")
            return []

    print("=" * 60)
    print("ПАРСЕР ЦППК (cpkrus.ru)")
    print("=" * 60)
    categories = get_categories()
    if not categories:
        print("Категории не найдены")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Обучающие программы"
    headers = ['Категория', 'Наименование программы', 'Объем программы',
               'Срок обучения', 'Получаемый документ', 'Требуемое образование',
               'Стоимость за весь курс', 'Ссылка на программу']
    ws.append(headers)
    row_num = 2
    for cat in categories:
        print(f'\nОбработка категории: {cat["name"]} ({cat["url"]})')
        courses = get_courses_from_category(cat['url'])
        print(f'  Найдено программ: {len(courses)}')
        for c in courses:
            ws.cell(row=row_num, column=1, value=cat['name'])
            ws.cell(row=row_num, column=2, value=c['title'])
            ws.cell(row=row_num, column=3, value=c['volume'])
            ws.cell(row=row_num, column=4, value=c['duration'])
            ws.cell(row=row_num, column=5, value=c['document'])
            ws.cell(row=row_num, column=6, value=c['education'])
            ws.cell(row=row_num, column=7, value=c['price_full'])
            ws.cell(row=row_num, column=8, value=c['url'])
            row_num += 1
        time.sleep(2)
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[col_letter].width = adjusted_width
    wb.save(OUTPUT_FILENAME)
    print(f'\nСохранено {row_num - 2} записей в {OUTPUT_FILENAME}')

# ========== ПАРСЕР №15: МЭИ ЮГА ==========

def run_parser_mei_yuga():
    """Парсер с сайта energetik-uc.narod.ru."""
    URL_TO_CATEGORY = {
        "elektro.html": "Электротехническое направление",
        "gpm.html": "Грузоподъемные машины и механизмы",
        "teplo.html": "Теплотехническое направление",
        "e-audit.html": "Техника и технология наземного транспорта",
        "kcn.html": "Охрана труда и пожарная безопасность",
        "business.html": "Бизнес образование",
        "prombez.html": "Промышленная безопасность",
    }
    PAGES_TO_PARSE = [f"https://energetik-uc.narod.ru/obuchen/{page}" for page in URL_TO_CATEGORY.keys()]
    OUTPUT_FILENAME = 'mei_yuga_data.xlsx'

    def get_soup(url):
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
        resp = requests.get(url, headers=headers, timeout=15)
        resp.raise_for_status()
        try:
            html = resp.content.decode('windows-1251')
        except UnicodeDecodeError:
            html = resp.content.decode('utf-8', errors='replace')
        return BeautifulSoup(html, 'html.parser')

    def extract_number(text):
        if not text:
            return None
        match = re.search(r'\d+', text)
        return int(match.group()) if match else None

    def get_category_from_row(row):
        cells = row.find_all('td')
        if not cells:
            return None
        first_cell = cells[0]
        colspan = int(first_cell.get('colspan', 1))
        if colspan >= 4:
            cat_text = first_cell.get_text(strip=True)
            if cat_text and any(kw in cat_text.lower() for kw in [
                'профессиональное обучение', 'дополнительное профессиональное образование',
                'дополнительное образование', 'повышение квалификации', 'профессиональная подготовка']):
                return cat_text
        return None

    def parse_educational_page(url, main_category):
        print(f"Парсинг: {url} -> {main_category}")
        soup = get_soup(url)
        if not soup:
            return []
        tables = soup.find_all('table')
        all_programs = []
        current_subcategory = ""
        for table in tables:
            rows = table.find_all('tr')
            if len(rows) < 3:
                continue
            header_row_idx = -1
            for idx, row in enumerate(rows[:4]):
                if any('Наименование' in cell.get_text(strip=True) for cell in row.find_all('td')):
                    header_row_idx = idx
                    break
            if header_row_idx == -1:
                continue
            raw_headers = []
            header_cells = rows[header_row_idx].find_all('td')
            for cell in header_cells:
                colspan = int(cell.get('colspan', 1))
                text = cell.get_text(strip=True).lower()
                if 'наименование' in text:
                    raw_headers.extend(['title'] * colspan)
                elif 'код' in text:
                    raw_headers.extend(['code'] * colspan)
                elif 'вид подготовки' in text or 'вид' in text:
                    raw_headers.extend(['training_type'] * colspan)
                elif 'дней' in text:
                    raw_headers.extend(['days'] * colspan)
                elif 'часов' in text and 'всего' not in text:
                    raw_headers.extend(['hours'] * colspan)
                elif 'всего' in text and 'часов' in text:
                    raw_headers.extend(['total_hours'] * colspan)
                else:
                    raw_headers.extend([f'extra_{len(raw_headers)}'] * colspan)
            while len(raw_headers) < 6:
                raw_headers.append(f'extra_{len(raw_headers)}')
            active_rowspans = {}
            for row in rows[header_row_idx + 1:]:
                subcat = get_category_from_row(row)
                if subcat:
                    current_subcategory = subcat
                    continue
                cells = row.find_all('td')
                if not cells:
                    continue
                row_data = []
                col_index = 0
                while col_index in active_rowspans:
                    row_data.append(active_rowspans[col_index]['text'])
                    active_rowspans[col_index]['remaining'] -= 1
                    if active_rowspans[col_index]['remaining'] == 0:
                        del active_rowspans[col_index]
                    col_index += 1
                for cell in cells:
                    while col_index in active_rowspans:
                        row_data.append(active_rowspans[col_index]['text'])
                        active_rowspans[col_index]['remaining'] -= 1
                        if active_rowspans[col_index]['remaining'] == 0:
                            del active_rowspans[col_index]
                        col_index += 1
                    cell_text = cell.get_text(strip=True)
                    rowspan = int(cell.get('rowspan', 1))
                    if rowspan > 1:
                        active_rowspans[col_index] = {'text': cell_text, 'remaining': rowspan - 1}
                    row_data.append(cell_text)
                    col_index += 1
                while col_index in active_rowspans:
                    row_data.append(active_rowspans[col_index]['text'])
                    active_rowspans[col_index]['remaining'] -= 1
                    if active_rowspans[col_index]['remaining'] == 0:
                        del active_rowspans[col_index]
                    col_index += 1
                while len(row_data) < len(raw_headers):
                    row_data.append('')
                program_raw = {raw_headers[i]: row_data[i].strip() for i in range(len(raw_headers))}
                title = program_raw.get('title', '')
                if not title or len(title) < 3:
                    continue
                unified = {
                    'Наименование программы': title,
                    'Код профессии': program_raw.get('code', ''),
                    'Вид подготовки': program_raw.get('training_type', ''),
                    'Категория': main_category,
                    'Подкатегория': current_subcategory,
                    'Дни (текст)': program_raw.get('days', ''),
                    'Часы (текст)': program_raw.get('hours', ''),
                    'Всего часов (текст)': program_raw.get('total_hours', ''),
                    'Источник': url
                }
                unified['Дни (число)'] = extract_number(unified['Дни (текст)'])
                unified['Часы (число)'] = extract_number(unified['Часы (текст)'])
                unified['Всего часов (число)'] = extract_number(unified['Всего часов (текст)'])
                all_programs.append(unified)
        return all_programs

    print("=" * 60)
    print("ПАРСЕР МЭИ ЮГА (energetik-uc.narod.ru)")
    print("=" * 60)
    all_records = []
    for url in PAGES_TO_PARSE:
        page_filename = url.split('/')[-1]
        main_category = URL_TO_CATEGORY.get(page_filename, "Другое")
        try:
            records = parse_educational_page(url, main_category)
            all_records.extend(records)
            print(f"  -> Найдено программ: {len(records)}")
        except Exception as e:
            print(f"Ошибка на {url}: {e}")
    if not all_records:
        print("Нет данных для сохранения")
        return
    df = pd.DataFrame(all_records)
    column_order = [
        'Наименование программы', 'Код профессии', 'Вид подготовки',
        'Категория', 'Подкатегория',
        'Дни (текст)', 'Часы (текст)', 'Всего часов (текст)',
        'Дни (число)', 'Часы (число)', 'Всего часов (число)', 'Источник'
    ]
    df = df[column_order]
    with pd.ExcelWriter(OUTPUT_FILENAME, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Все программы', index=False)
        worksheet = writer.sheets['Все программы']
        for column in worksheet.columns:
            max_len = 0
            col_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_len:
                        max_len = len(str(cell.value))
                except Exception:
                    pass
            worksheet.column_dimensions[col_letter].width = min(max_len + 2, 60)
    print(f"Сохранено {OUTPUT_FILENAME} ({len(df)} записей)")

# ========== ПАРСЕР №16: УПК ==========

def run_parser_upk():
    """Парсер с сайта с кириллическим доменом (УПК)."""
    from openpyxl import Workbook
    BASE_URL = "https://xn--21-1lctt.xn--p1ai"
    START_URL = urljoin(BASE_URL, "/specialties.php")
    OUTPUT_FILE = 'upk_data.xlsx'
    SAVE_EVERY = 20
    HEADERS = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    }

    def get_html(url):
        try:
            r = requests.get(url, headers=HEADERS, timeout=10)
            r.raise_for_status()
            r.encoding = 'utf-8'
            return r.text
        except Exception as e:
            print(f"  Ошибка загрузки: {e}")
            return None

    def extract_specialty_from_html(html):
        script_match = re.search(r'<script\s+type="module">(.*?)</script>', html, re.DOTALL)
        if not script_match:
            return None
        script_content = script_match.group(1)
        pos = script_content.find('specialty:')
        if pos == -1:
            return None
        brace_start = script_content.find('{', pos)
        if brace_start == -1:
            return None
        brace_count = 0
        i = brace_start
        in_string = False
        escape = False
        while i < len(script_content):
            ch = script_content[i]
            if not escape:
                if ch == '"' and not in_string:
                    in_string = True
                elif ch == '"' and in_string:
                    in_string = False
                elif ch == '\\' and in_string:
                    escape = True
                elif not in_string:
                    if ch == '{':
                        brace_count += 1
                    elif ch == '}':
                        brace_count -= 1
                        if brace_count == 0:
                            break
            else:
                escape = False
            i += 1
        if brace_count != 0:
            return None
        json_str = script_content[brace_start:i + 1]
        json_str = json_str.replace(':None', ':null')
        try:
            return json.loads(json_str)
        except json.JSONDecodeError as e:
            print(f"  Ошибка JSON: {e}")
            return None

    print("=" * 60)
    print("ПАРСЕР УПК")
    print("=" * 60)
    html = get_html(START_URL)
    if not html:
        return
    soup = BeautifulSoup(html, 'lxml')
    accordions = soup.find_all('div', class_='accordion-item')
    print(f"Найдено категорий: {len(accordions)}")
    all_programs = []
    for accordion in accordions:
        title_elem = accordion.find('div', class_='accordion-title')
        category = title_elem.get_text(strip=True) if title_elem else "Без категории"
        prog_list = accordion.find('div', class_='sub-programs-list')
        if not prog_list:
            continue
        for a in prog_list.find_all('a', class_='sub-program-item', href=True):
            code_elem = a.find('span', class_='sub-program-code')
            name_elem = a.find('span', class_='sub-program-name')
            code = code_elem.get_text(strip=True) if code_elem else ""
            name = name_elem.get_text(strip=True) if name_elem else ""
            url = urljoin(BASE_URL, a['href'])
            all_programs.append((category, code, name, url))
    print(f"Всего программ: {len(all_programs)}")
    headers = [
        "Категория", "Код программы", "Наименование", "Ссылка",
        "Общая длительность (часы)", "Теория (часы)", "Практика (часы)",
        "Стоимость очного обучения", "Стоимость дистанционного обучения"
    ]
    results = []
    for idx, (cat, code, name, url) in enumerate(all_programs, start=1):
        print(f"[{idx}/{len(all_programs)}] {code} - {name}")
        page_html = get_html(url)
        if not page_html:
            results.append([cat, code, name, url, "", "", "", "", ""])
            continue
        spec = extract_specialty_from_html(page_html)
        if not spec:
            results.append([cat, code, name, url, "", "", "", "", ""])
            continue
        total_hours = spec.get('hours', '')
        theory = spec.get('theory_learning', '')
        practice = spec.get('practice_learning', '')
        cost_offline = spec.get('cost', '')
        cost_online = spec.get('cost_distant', '')
        results.append([cat, code, name, url, total_hours, theory, practice, cost_offline, cost_online])
        if idx % SAVE_EVERY == 0 or idx == len(all_programs):
            wb = Workbook()
            ws = wb.active
            ws.append(headers)
            for row in results:
                ws.append(row)
            wb.save(OUTPUT_FILE)
            print(f"  >>> Сохранено {idx} записей в {OUTPUT_FILE}")
        time.sleep(0.3)
    print(f"\nОбработано {len(results)} программ. Данные сохранены в {OUTPUT_FILE}")

# ========== ПАРСЕР №17: БУРОВИК ==========

def run_parser_burovik():
    """Парсер с сайта уц-буровик.рф."""
    URL = "https://уц-буровик.рф/курсы/"
    HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
    OUTPUT_FILENAME = 'burovik_data.xlsx'

    def extract_hours(duration_text):
        preparation = retraining = advanced = ''
        prep_match = re.search(r'подготовка\s+(\d+)\s*час', duration_text, re.IGNORECASE)
        if prep_match:
            preparation = prep_match.group(1)
        retrain_match = re.search(r'переподготовка\s+(\d+)\s*час', duration_text, re.IGNORECASE)
        if retrain_match:
            retraining = retrain_match.group(1)
        adv_match = re.search(r'повышение\s+квалификации\s+(\d+)\s*час', duration_text, re.IGNORECASE)
        if adv_match:
            advanced = adv_match.group(1)
        return preparation, retraining, advanced

    def parse_courses():
        try:
            response = requests.get(URL, headers=HEADERS, timeout=15, verify=False)
            response.raise_for_status()
            response.encoding = 'utf-8'
        except Exception as e:
            print(f"Ошибка загрузки: {e}")
            return []
        soup = BeautifulSoup(response.text, 'html.parser')
        spoilers = soup.find_all('div', class_='su-spoiler-content')
        if not spoilers:
            print("Не найдено блоков .su-spoiler-content")
            return []
        courses_data = []
        for spoiler in spoilers:
            items = spoiler.find_all('li', recursive=True)
            for li in items:
                title_tag = li.find('strong')
                if title_tag:
                    course_name = title_tag.get_text(strip=True)
                else:
                    course_name = li.get_text(strip=True).split('Форма обучения')[0].strip()
                table = li.find('div', class_='su-table')
                if not table:
                    table = li.find('table')
                if not table:
                    continue
                rows = table.find_all('tr')
                if not rows:
                    continue
                first_row_cells = rows[0].find_all('td')
                if len(first_row_cells) == 3 and \
                   any('Форма' in cell.get_text() for cell in first_row_cells) and \
                   any('Срок' in cell.get_text() for cell in first_row_cells) and \
                   any('Стоимость' in cell.get_text() for cell in first_row_cells):
                    if len(rows) >= 2:
                        value_cells = rows[1].find_all('td')
                        if len(value_cells) == 3:
                            form = value_cells[0].get_text(strip=True)
                            duration = value_cells[1].get_text(strip=True)
                            cost = value_cells[2].get_text(strip=True)
                        else:
                            form = duration = cost = ''
                    else:
                        form = duration = cost = ''
                else:
                    form = duration = cost = ''
                    for row in rows:
                        cells = row.find_all('td')
                        if len(cells) < 2:
                            continue
                        label = cells[0].get_text(strip=True)
                        value = cells[1].get_text(strip=True)
                        if 'Форма обучения' in label:
                            form = value
                        elif 'Срок обучения' in label:
                            duration = value
                        elif 'Стоимость обучения' in label:
                            cost = value
                prep, retrain, advanced = extract_hours(duration)
                courses_data.append({
                    'Наименование программы': course_name, 'Форма обучения': form,
                    'Срок обучения': duration, 'Подготовка (часы)': prep,
                    'Переподготовка (часы)': retrain, 'Повышение квалификации (часы)': advanced,
                    'Стоимость обучения': cost
                })
        return courses_data

    print("=" * 60)
    print("ПАРСЕР БУРОВИК")
    print("=" * 60)
    courses = parse_courses()
    if courses:
        print(f"Найдено программ: {len(courses)}")
        df = pd.DataFrame(courses)
        with pd.ExcelWriter(OUTPUT_FILENAME, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Курсы')
            worksheet = writer.sheets['Курсы']
            for column in worksheet.columns:
                max_len = 0
                col_letter = column[0].column_letter
                for cell in column:
                    try:
                        max_len = max(max_len, len(str(cell.value)))
                    except Exception:
                        pass
                worksheet.column_dimensions[col_letter].width = min(max_len + 2, 50)
        print(f"Сохранено {len(courses)} записей в {OUTPUT_FILENAME}")
    else:
        print("Не удалось извлечь данные")

# ========== ПАРСЕР №18: ПАТРИОТ ==========

def run_parser_patriot():
    """Парсер с сайта stc-patriot.ru."""
    OUTPUT_FILENAME = 'patriot_data.xlsx'

    def fetch_page_content(url):
        try:
            response = requests.get(url)
            response.raise_for_status()
            response.encoding = 'utf-8'
            return response.text
        except requests.RequestException as e:
            print(f"Ошибка загрузки: {e}")
            return None

    def parse_html_tables(html_content, max_tables=4):
        soup = BeautifulSoup(html_content, 'html.parser')
        all_tables = soup.find_all('table')
        if not all_tables:
            print("HTML-таблицы не найдены")
            return []
        print(f"Найдено HTML-таблиц: {len(all_tables)}")
        dataframes = []
        for i, table in enumerate(all_tables[:max_tables]):
            try:
                rows = []
                for row in table.find_all('tr'):
                    cells = row.find_all(['td', 'th'])
                    row_data = [cell.get_text(strip=True) for cell in cells]
                    if row_data:
                        rows.append(row_data)
                if not rows:
                    continue
                if len(rows) > 1:
                    df = pd.DataFrame(rows[1:], columns=rows[0])
                else:
                    df = pd.DataFrame(rows)
                dataframes.append(df)
                print(f"  Таблица {i + 1}: {df.shape[0]} строк, {df.shape[1]} столбцов")
            except Exception as e:
                print(f"  Ошибка в таблице {i + 1}: {e}")
        return dataframes

    def combine_tables(dataframes):
        if not dataframes:
            return pd.DataFrame()
        if len(dataframes) == 1:
            return dataframes[0]
        all_columns = set()
        for df in dataframes:
            all_columns.update(df.columns)
        all_columns = sorted(list(all_columns))
        aligned_dfs = []
        for df in dataframes:
            for col in all_columns:
                if col not in df.columns:
                    df[col] = pd.NA
            aligned_dfs.append(df[all_columns])
        combined = pd.concat(aligned_dfs, ignore_index=True)
        return combined

    print("=" * 60)
    print("ПАРСЕР ПАТРИОТ (stc-patriot.ru)")
    print("=" * 60)
    url = "https://stc-patriot.ru/prajs-list/"
    print(f"Загрузка страницы: {url}")
    html = fetch_page_content(url)
    if not html:
        return
    tables = parse_html_tables(html, max_tables=4)
    if not tables:
        return
    combined_df = combine_tables(tables)
    with pd.ExcelWriter(OUTPUT_FILENAME, engine='openpyxl') as writer:
        combined_df.to_excel(writer, sheet_name="Все программы", index=False)
    print(f"Сохранено в {OUTPUT_FILENAME}, всего строк: {len(combined_df)}")

# ========== ПАРСЕР №19: ЧУВАШСКИЙ УЧЕБНО-КУРСОВОЙ КОМБИНАТ ==========

def run_parser_chuvash():
    """Парсер с сайта chukk.ru."""
    OUTPUT_FILENAME = 'chuvash_data.xlsx'

    def parse_educational_programs(url):
        response = requests.get(url, verify=False)
        response.encoding = 'windows-1251'
        soup = BeautifulSoup(response.text, 'html.parser')
        a_a_div = soup.find('div', class_='a_a')
        if not a_a_div:
            raise ValueError("Не найден контейнер с классом 'a_a'")
        text_blocks = a_a_div.find_all('div', class_='text')
        all_programs = []
        for block in text_blocks:
            tables = block.find_all('table')
            for table in tables:
                rows = table.find_all('tr')
                for row in rows:
                    cells = row.find_all('td')
                    if not cells:
                        continue
                    program = {}
                    if len(cells) >= 3:
                        program['№ п/п'] = cells[0].get_text(strip=True) if len(cells) > 0 else ''
                        program['Код'] = cells[1].get_text(strip=True) if len(cells) > 1 else ''
                        program['Наименование'] = cells[2].get_text(strip=True) if len(cells) > 2 else ''
                        if len(cells) > 5:
                            program['Теоретическое обучение (час)'] = cells[3].get_text(strip=True) if len(cells) > 3 else ''
                            program['Производственное обучение (час)'] = cells[4].get_text(strip=True) if len(cells) > 4 else ''
                            program['Общая продолжительность (час)'] = cells[5].get_text(strip=True) if len(cells) > 5 else ''
                        if len(cells) > 6:
                            program['Стоимость очного обучения (руб)'] = cells[6].get_text(strip=True) if len(cells) > 6 else ''
                        if len(cells) > 7:
                            program['Стоимость электронного обучения (руб)'] = cells[7].get_text(strip=True) if len(cells) > 7 else ''
                        if program['Наименование']:
                            all_programs.append(program)
        return all_programs

    print("=" * 60)
    print("ПАРСЕР ЧУВАШСКИЙ УЧЕБНО-КУРСОВОЙ КОМБИНАТ (chukk.ru)")
    print("=" * 60)
    url = "https://chukk.ru/perechen-obrazovatelnykh-uslug"
    try:
        programs = parse_educational_programs(url)
        if programs:
            df = pd.DataFrame(programs)
            column_order = ['№ п/п', 'Код', 'Наименование',
                            'Теоретическое обучение (час)', 'Производственное обучение (час)',
                            'Общая продолжительность (час)',
                            'Стоимость очного обучения (руб)', 'Стоимость электронного обучения (руб)']
            existing_columns = [col for col in column_order if col in df.columns]
            df = df[existing_columns]
            with pd.ExcelWriter(OUTPUT_FILENAME, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Образовательные программы', index=False)
                worksheet = writer.sheets['Образовательные программы']
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except Exception:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
            print(f"Данные сохранены в {OUTPUT_FILENAME}. Всего программ: {len(programs)}")
        else:
            print("Не удалось извлечь данные")
    except Exception as e:
        print(f"Произошла ошибка: {e}")

# ========== ПАРСЕР №20: МРСК УРАЛА ==========

def run_parser_mrsk_ural():
    """Парсер скачивает готовый Excel с сайта uc-mrsk-ural.ru."""
    OUTPUT_FILENAME = 'mrsk_ural_data.xlsx'
    file_url = "https://uc-mrsk-ural.ru/media/files/perechen_programm_na_2026.xlsx"
    print("=" * 60)
    print("ПАРСЕР МРСК УРАЛА (uc-mrsk-ural.ru)")
    print("=" * 60)
    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
    try:
        print(f"Загрузка из: {file_url}")
        with requests.get(file_url, stream=True, headers=headers, verify=False) as response:
            response.raise_for_status()
            total_size = int(response.headers.get('content-length', 0))
            downloaded = 0
            with open(OUTPUT_FILENAME, 'wb') as file:
                for chunk in response.iter_content(chunk_size=8192):
                    if chunk:
                        file.write(chunk)
                        downloaded += len(chunk)
                        if total_size:
                            percent = (downloaded / total_size) * 100
                            print(f"\rПрогресс: {percent:.1f}% ({downloaded} / {total_size} байт)", end='')
        print(f"\nФайл сохранён: {OUTPUT_FILENAME}")
    except requests.exceptions.RequestException as e:
        print(f"\nОшибка загрузки: {e}")

# ========== ПАРСЕР №21: ИГЭУ ==========

def run_parser_igeu():
    """Парсер PDF с сайта ispu.ru (ИГЭУ)."""
    pdf_urls = [
        "http://ispu.ru/sites/default/files/2025-06/ПЛАН%202026%20генерирующие%20компании.pdf",
        "http://ispu.ru/sites/default/files/2025-06/ПЛАН%202026%20промышленные%20предприятия.pdf",
        "http://ispu.ru/sites/default/files/2025-06/ПЛАН%202026%20сетевые%20компании.pdf"
    ]
    source_names = ["Генерирующие_компании", "Промышленные_предприятия", "Сетевые_компании"]
    output_file = 'igeu_data.xlsx'
    output_sheet = "Все_программы"

    def safe_str(value):
        if pd.isna(value):
            return ''
        return str(value).strip()

    def merge_program_rows(df):
        if df.empty:
            return df
        df = df.fillna('').astype(str).replace('nan', '')
        new_rows = []
        current_row = None
        for idx in range(len(df)):
            first_cell = df.iloc[idx, 0].strip()
            is_new_program = bool(re.match(r'^\d+\.', first_cell))
            if is_new_program:
                if current_row is not None:
                    new_rows.append(current_row)
                current_row = df.iloc[idx].tolist()
            else:
                if current_row is None:
                    current_row = df.iloc[idx].tolist()
                else:
                    for col in range(df.shape[1]):
                        curr_val = df.iloc[idx, col]
                        if curr_val:
                            current_row[col] = f"{current_row[col]} {curr_val}".strip()
        if current_row is not None:
            new_rows.append(current_row)
        return pd.DataFrame(new_rows)

    def extract_hours_cost_duration(df):
        hours_pat = re.compile(r'(\d+(?:[.,]\d+)?)\s*(?:ак\.?\s*час|ак\.?\s*ч|час(?:ов)?|ч\.?)', re.I)
        cost_pat = re.compile(r'(\d+(?:[.,]\d+)?)\s*(?:руб|р\.?|₽)', re.I)
        duration_pat = re.compile(r'(\d+(?:[.,]\d+)?)\s*(?:год(?:а|ов)?|г\.|мес(?:яц(?:ев)?)?|курс(?:а|ов)?)', re.I)
        def find_first(text, pattern):
            if not isinstance(text, str):
                text = safe_str(text)
            m = pattern.search(text)
            return m.group(1).replace(',', '.') if m else ''
        df['Часы'] = ''
        df['Стоимость'] = ''
        df['Срок'] = ''
        for idx in range(len(df)):
            row_text = ' '.join([safe_str(df.iloc[idx, c]) for c in range(df.shape[1] - 3)])
            df.at[idx, 'Часы'] = find_first(row_text, hours_pat)
            df.at[idx, 'Стоимость'] = find_first(row_text, cost_pat)
            df.at[idx, 'Срок'] = find_first(row_text, duration_pat)
        return df

    def process_pdf_from_url(url, source_name):
        print(f"Начинаю обработку: {source_name}")
        try:
            response = requests.get(url, verify=False, timeout=30)
            response.raise_for_status()
            pdf_file = io.BytesIO(response.content)
            tables = camelot.read_pdf(pdf_file, flavor='stream', pages='all', row_tol=10, split_text=False, edge_tol=500)
            if not tables:
                print(f"  Предупреждение: В PDF '{source_name}' таблицы не найдены")
                return pd.DataFrame()
            print(f"  Найдено таблиц: {len(tables)}")
            combined_df = pd.concat([table.df for table in tables], ignore_index=True)
            combined_df = combined_df.fillna('')
            combined_df = combined_df.replace(r'^\s*$', '', regex=True)
            combined_df = combined_df[~combined_df.apply(lambda row: row.astype(str).str.strip().eq('').all(), axis=1)]
            combined_df = merge_program_rows(combined_df)
            combined_df = extract_hours_cost_duration(combined_df)
            combined_df.insert(0, 'Источник', source_name)
            print(f"  Обработка '{source_name}' завершена. Получено строк: {len(combined_df)}")
            return combined_df
        except Exception as e:
            print(f"  Ошибка при обработке {source_name}: {e}")
            return pd.DataFrame()

    print("=" * 60)
    print("ПАРСЕР ИГЭУ (ispu.ru)")
    print("=" * 60)
    all_data = pd.DataFrame()
    for url, src in zip(pdf_urls, source_names):
        df = process_pdf_from_url(url, src)
        if not df.empty:
            all_data = pd.concat([all_data, df], ignore_index=True)
    if not all_data.empty:
        print(f"\nВсего собрано записей: {len(all_data)}")
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            all_data.to_excel(writer, sheet_name=output_sheet, index=False, header=True)
            worksheet = writer.sheets[output_sheet]
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                worksheet.column_dimensions[column_letter].width = min(max_length + 2, 60)
        print(f"Файл сохранён: {os.path.abspath(output_file)}")
    else:
        print("Нет данных для сохранения")

# ========== ПАРСЕР №22: ЦПО ==========

def run_parser_cpo():
    """Парсер с сайта educpo.ru."""
    OUTPUT_FILENAME = 'cpo_data.xlsx'

    def parse_qualification(soup):
        data = []
        for element in soup.select("#content .span12"):
            for h2 in element.find_all("h2"):
                category = h2.get_text(strip=True).replace("## ", "")
                next_table = h2.find_next_sibling("div", class_="pricing-tables")
                if next_table:
                    rows = next_table.find_all("tr")
                    for row in rows[1:]:
                        cols = row.find_all("td")
                        if len(cols) >= 4:
                            data.append({
                                "Категория": category,
                                "Номер п/п": cols[0].get_text(strip=True),
                                "Наименование программы": cols[1].get_text(strip=True),
                                "Количество часов": cols[2].get_text(strip=True),
                                "Код программы": cols[3].get_text(strip=True)
                            })
        return data

    def parse_retraining(soup):
        data = []
        for group in soup.select(".accordion-group"):
            toggle = group.select_one(".accordion-toggle")
            if not toggle:
                continue
            raw_text = toggle.get_text(strip=True)
            if " — " in raw_text:
                parts = raw_text.rsplit(" — ", 1)
                name = parts[0].strip()
                hours_raw = parts[1].strip()
            elif " - " in raw_text:
                parts = raw_text.rsplit(" - ", 1)
                name = parts[0].strip()
                hours_raw = parts[1].strip()
            else:
                name = raw_text
                hours_raw = ""
            if hours_raw:
                hours_match = re.search(r'^[\d/]+', hours_raw)
                hours = hours_match.group(0) if hours_match else hours_raw
            else:
                hours = ""
            data.append({
                "Категория": "Профессиональная переподготовка",
                "Номер п/п": "",
                "Наименование программы": name,
                "Количество часов": hours,
                "Код программы": ""
            })
        return data

    def fetch_soup(url):
        headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
        response = requests.get(url, headers=headers, verify=False, timeout=10)
        response.encoding = "utf-8"
        response.raise_for_status()
        return BeautifulSoup(response.text, "html.parser")

    print("=" * 60)
    print("ПАРСЕР ЦПО (educpo.ru)")
    print("=" * 60)
    urls = {
        "qualification": "https://educpo.ru/dopolnitelnoe-obrazovanie/povyshenie-kvalifikacii",
        "retraining": "https://educpo.ru/dopolnitelnoe-obrazovanie/professionalnaya-perepodgotovka"
    }
    all_data = []
    try:
        soup = fetch_soup(urls["qualification"])
        qual_data = parse_qualification(soup)
        all_data.extend(qual_data)
        print(f"Повышение квалификации: {len(qual_data)} программ")
    except Exception as e:
        print(f"Ошибка при парсинге повышения квалификации: {e}")
    try:
        soup = fetch_soup(urls["retraining"])
        ret_data = parse_retraining(soup)
        all_data.extend(ret_data)
        print(f"Профессиональная переподготовка: {len(ret_data)} программ")
    except Exception as e:
        print(f"Ошибка при парсинге переподготовки: {e}")
    if not all_data:
        print("Нет данных для сохранения")
        return
    df = pd.DataFrame(all_data)
    df.to_excel(OUTPUT_FILENAME, index=False, engine="openpyxl")
    print(f"Сохранено всего {len(all_data)} записей в {OUTPUT_FILENAME}")

# ========== ПАРСЕР №23: ЮРГПУ(НПИ) ==========

def run_parser_yurgpu():
    """Парсер с сайта npi-tu.ru."""
    BASE_URL = 'https://npi-tu.ru'
    HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
    OUTPUT_FILENAME = 'yurgpu_data.xlsx'

    def get_soup(url):
        try:
            response = requests.get(url, headers=HEADERS, timeout=15, verify=False)
            response.raise_for_status()
            return BeautifulSoup(response.text, 'html.parser')
        except requests.exceptions.RequestException as e:
            print(f'Ошибка при загрузке {url}: {e}')
            return None

    def get_category_links(main_url):
        soup = get_soup(main_url)
        if not soup:
            return []
        main_block = soup.find('main', class_='sf-main-area')
        if not main_block:
            return []
        ul = main_block.find('ul')
        if not ul:
            return []
        links = []
        for li in ul.find_all('li'):
            a = li.find('a')
            if a and a.get('href'):
                href = a['href']
                full_url = urljoin(BASE_URL, href)
                category_name = a.get_text(strip=True)
                links.append((category_name, full_url))
        return links

    def parse_program_page(category_name, page_url):
        soup = get_soup(page_url)
        if not soup:
            return []
        programs = []
        table = soup.find('table')
        if table:
            rows = table.find_all('tr')
            for row in rows[1:]:
                cols = row.find_all('td')
                if len(cols) >= 3:
                    name = cols[0].get_text(strip=True)
                    hours = cols[1].get_text(strip=True)
                    price = cols[2].get_text(strip=True)
                    programs.append({'category': category_name, 'link': page_url,
                                     'name': name, 'hours': hours, 'price': price})
            if programs:
                return programs
        items = soup.find_all('div', class_='program-item')
        if items:
            for item in items:
                name_elem = item.find(class_='program-name')
                hours_elem = item.find(class_='program-hours')
                price_elem = item.find(class_='program-price')
                if name_elem:
                    programs.append({
                        'category': category_name, 'link': page_url,
                        'name': name_elem.get_text(strip=True),
                        'hours': hours_elem.get_text(strip=True) if hours_elem else '',
                        'price': price_elem.get_text(strip=True) if price_elem else ''
                    })
            if programs:
                return programs
        return []

    print("=" * 60)
    print("ПАРСЕР ЮРГПУ(НПИ) (npi-tu.ru)")
    print("=" * 60)
    main_url = 'https://npi-tu.ru/university/faculty/ino/tsdo/povyshenie-kvalifikatsii/ds/'
    categories = get_category_links(main_url)
    print(f'Найдено категорий: {len(categories)}')
    if not categories:
        return
    all_data = []
    for idx, (cat_name, cat_url) in enumerate(categories, start=1):
        print(f'[{idx}/{len(categories)}] Парсинг: {cat_name} ({cat_url})')
        programs = parse_program_page(cat_name, cat_url)
        if programs:
            all_data.extend(programs)
            print(f'  -> Найдено {len(programs)} программ')
        time.sleep(1)
    if all_data:
        df = pd.DataFrame(all_data)
        df = df[['category', 'link', 'name', 'hours', 'price']]
        df.to_excel(OUTPUT_FILENAME, index=False, engine='openpyxl')
        print(f'Сохранено {len(df)} записей в {OUTPUT_FILENAME}')
    else:
        print('Нет данных для сохранения')

# ========== ПАРСЕР №24: ИПО ==========

def run_parser_ipo():
    """Парсер с сайта ipo.msk.ru."""
    BASE_URL = 'https://ipo.msk.ru'
    HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'}
    SECTIONS = [
        {'name': 'Профессиональная переподготовка', 'url': '/professionalnaja-perepodgotovka/'},
        {'name': 'Повышение квалификации', 'url': '/povyshenie-kvalifikacii/'}
    ]
    OUTPUT_FILENAME = 'ipo_data.xlsx'

    def get_soup(url):
        try:
            resp = requests.get(url, headers=HEADERS, timeout=15, verify=False)
            resp.raise_for_status()
            return BeautifulSoup(resp.text, 'html.parser')
        except Exception as e:
            print(f'Ошибка загрузки {url}: {e}')
            return None

    def get_category_links(section_url):
        full_url = urljoin(BASE_URL, section_url)
        soup = get_soup(full_url)
        if not soup:
            return []
        ul = soup.select_one('div.list-prog-dpo ul.nily-field-items')
        if not ul:
            return []
        categories = []
        for li in ul.find_all('li', class_='nily-field-items__nily-item'):
            a = li.find('a', class_='nily-item__nily-link')
            if a and a.get('href'):
                href = a['href']
                full_cat_url = urljoin(BASE_URL, href)
                name_elem = a.find('h3', class_='nily-link__program-name')
                cat_name = name_elem.get_text(strip=True) if name_elem else 'Без названия'
                categories.append({'name': cat_name, 'url': full_cat_url})
        return categories

    def get_program_links(category_url):
        soup = get_soup(category_url)
        if not soup:
            return []
        program_urls = []
        list_dpo = soup.find('div', class_='list-prog-dpo')
        if list_dpo:
            ul = list_dpo.find('ul')
            if ul:
                for li in ul.find_all('li'):
                    a = li.find('a', href=True)
                    if a:
                        full_url = urljoin(BASE_URL, a['href'])
                        program_urls.append(full_url)
                if program_urls:
                    return list(set(program_urls))
        program_list = soup.find('div', class_='program-list')
        if program_list:
            for a in program_list.find_all('a', class_='program-list__program-title', href=True):
                full_url = urljoin(BASE_URL, a['href'])
                program_urls.append(full_url)
            if program_urls:
                return list(set(program_urls))
        return program_urls

    def parse_program_page(program_url):
        soup = get_soup(program_url)
        if not soup:
            return {}
        data = {'url': program_url}
        h1 = soup.find('h1')
        data['name'] = h1.get_text(strip=True) if h1 else ''
        data['type'] = ''
        data['format'] = ''
        promo = soup.find('div', class_='promo-block')
        if promo:
            list_blocks = promo.find_all('div', class_='promo-block-item-list')
            hours = ''
            duration = ''
            for block in list_blocks:
                text = block.get_text(separator=' ', strip=True)
                match = re.search(r'(\d+)\s*(?:час|ч\.)', text, re.IGNORECASE)
                if match:
                    hours = match.group(1) + ' ч.'
                dur_match = re.search(r'до\s*(\d+)\s*месяц', text, re.IGNORECASE)
                if dur_match:
                    duration = f"до {dur_match.group(1)} месяца"
            data['hours'] = hours
            data['duration'] = duration
            price_new = promo.find('div', class_='promo-block-item-pnew')
            data['price_new'] = price_new.get_text(strip=True) if price_new else ''
            price_old = promo.find('div', class_='promo-block-item-pold')
            data['price_old'] = price_old.get_text(strip=True) if price_old else ''
        else:
            data['hours'] = ''
            data['duration'] = ''
            data['price_new'] = ''
            data['price_old'] = ''
        return data

    print("=" * 60)
    print("ПАРСЕР ИПО (ipo.msk.ru)")
    print("=" * 60)
    all_data = []
    for section in SECTIONS:
        print(f'Обработка раздела: {section["name"]}')
        categories = get_category_links(section['url'])
        print(f'  Найдено категорий: {len(categories)}')
        for cat in categories:
            print(f'    Категория: {cat["name"]} ({cat["url"]})')
            program_links = get_program_links(cat['url'])
            print(f'      Найдено программ: {len(program_links)}')
            for p_url in program_links:
                prog_data = parse_program_page(p_url)
                if prog_data and prog_data.get('name'):
                    prog_data['section'] = section['name']
                    prog_data['category'] = cat['name']
                    all_data.append(prog_data)
                time.sleep(0.5)
    if all_data:
        df = pd.DataFrame(all_data)
        columns = ['section', 'category', 'name', 'type', 'format', 'hours', 'duration', 'price_new', 'price_old', 'url']
        df = df[columns]
        df.to_excel(OUTPUT_FILENAME, index=False, engine='openpyxl')
        print(f'Сохранено {len(df)} записей в {OUTPUT_FILENAME}')
    else:
        print('Нет данных')

# ========== ПАРСЕР №25: ЦЕНТР ЛИЧНОСТНОГО РОСТА ==========

def run_parser_lichnostniy_rost():
    """Парсер с сайта lichtnostniyrost.ru."""
    BASE_URL = "https://lichtnostniyrost.ru"
    START_URL = urljoin(BASE_URL, "/online-education/")
    HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
    OUTPUT_FILENAME = 'lichnostniy_rost_data.xlsx'

    def get_soup(url, retries=3):
        for attempt in range(retries):
            try:
                response = requests.get(url, headers=HEADERS, verify=False, timeout=10)
                response.encoding = 'windows-1251'
                return BeautifulSoup(response.text, 'html.parser')
            except Exception as e:
                print(f"Попытка {attempt + 1}/{retries} загрузить {url} не удалась: {e}")
                time.sleep(2)
        return None

    def parse_program_detail(program_url, default_name, default_hours):
        soup = get_soup(program_url)
        if not soup:
            return {'name': default_name, 'for_whom': '', 'study_form': '', 'hours': default_hours, 'price': ''}
        for_whom = study_form = hours = ''
        properties = soup.select('.catalog-detail__property')
        for prop in properties:
            title_elem = prop.select_one('.catalog-detail__property-title')
            value_elem = prop.select_one('.catalog-detail__property-value')
            if title_elem and value_elem:
                title = title_elem.get_text(strip=True)
                value = value_elem.get_text(strip=True)
                if 'Для кого' in title:
                    for_whom = value
                elif 'Форма обучения' in title:
                    study_form = value
                elif 'Объем ак. часов' in title:
                    hours = value
        price_elem = soup.select_one('.catalog-detail__price')
        price = price_elem.get_text(strip=True) if price_elem else ''
        price = re.sub(r'руб\.?', '', price).strip()
        title_elem = soup.find('h1')
        name = title_elem.get_text(strip=True) if title_elem else default_name
        return {'name': name, 'for_whom': for_whom, 'study_form': study_form, 'hours': hours, 'price': price}

    def parse_programs_from_list(soup, category_name, subcategory_name=''):
        programs = []
        items = soup.select('.catalog-list__item')
        for item in items:
            link_elem = item.select_one('.catalog-list__link')
            if not link_elem:
                continue
            href = link_elem.get('href')
            if not href:
                continue
            program_url = urljoin(BASE_URL, href)
            title_elem = item.select_one('.catalog-list__title')
            title_text = title_elem.get_text(strip=True) if title_elem else ''
            name_match = re.search(r'«(.*?)»', title_text)
            name = name_match.group(1) if name_match else title_text
            hours_match = re.search(r'(\d+)\s*ак\.?часов?', title_text)
            hours = hours_match.group(1) if hours_match else ''
            if not hours:
                length_elem = item.select_one('.catalog-list__length')
                if length_elem:
                    hours = length_elem.get_text(strip=True)
            detail = parse_program_detail(program_url, name, hours)
            programs.append({
                'Наименование': detail['name'], 'Категория': category_name,
                'Подкатегория': subcategory_name, 'Ссылка на программу': program_url,
                'Для кого': detail['for_whom'], 'Форма обучения': detail['study_form'],
                'Объем ак. часов': detail['hours'], 'Стоимость': detail['price']
            })
            time.sleep(0.5)
        return programs

    def parse_category(category_url, category_name):
        soup = get_soup(category_url)
        if not soup:
            return []
        all_programs = []
        subcategory_links = soup.select('a.section-list__link_depth_more')
        if subcategory_links:
            for link in subcategory_links:
                sub_name = link.get_text(strip=True)
                sub_url = urljoin(BASE_URL, link.get('href'))
                print(f"  Обработка подкатегории: {sub_name}")
                sub_soup = get_soup(sub_url)
                if sub_soup:
                    programs = parse_programs_from_list(sub_soup, category_name, sub_name)
                    all_programs.extend(programs)
                    print(f"    Найдено программ: {len(programs)}")
        else:
            programs = parse_programs_from_list(soup, category_name, '')
            all_programs.extend(programs)
            print(f"    Найдено программ: {len(programs)}")
        return all_programs

    print("=" * 60)
    print("ПАРСЕР ЦЕНТР ЛИЧНОСТНОГО РОСТА")
    print("=" * 60)
    main_soup = get_soup(START_URL)
    if not main_soup:
        return
    category_links = main_soup.select('a.section-list__link_depth_1')
    all_data = []
    for link in category_links:
        cat_name = link.get_text(strip=True)
        cat_url = urljoin(BASE_URL, link.get('href'))
        print(f"Обработка категории: {cat_name}")
        programs = parse_category(cat_url, cat_name)
        all_data.extend(programs)
        print(f"  Всего в категории: {len(programs)} программ\n")
    if all_data:
        df = pd.DataFrame(all_data)
        df.to_excel(OUTPUT_FILENAME, index=False)
        print(f"Собрано {len(all_data)} программ. Данные сохранены в {OUTPUT_FILENAME}")
    else:
        print("Не удалось найти ни одной программы")

# ========== ПАРСЕР №26: ЭНЕРГЕТИК НН ==========

def run_parser_energetik_nn():
    """Парсер с сайта energetik-nn.ru."""
    BASE_URL = "https://energetik-nn.ru"
    START_URL = BASE_URL + "/"
    HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
    OUTPUT_FILENAME = 'energetik_nn_data.xlsx'

    def get_soup(url):
        try:
            response = requests.get(url, headers=HEADERS, verify=False, timeout=10)
            response.encoding = 'utf-8'
            return BeautifulSoup(response.text, 'html.parser')
        except Exception as e:
            print(f"  Ошибка загрузки {url}: {e}")
            return None

    def extract_number(text):
        match = re.search(r'(\d+)', text)
        return match.group(1) if match else ''

    def parse_detail_page(url):
        soup = get_soup(url)
        if not soup:
            return {'price': '', 'format': ''}
        price = ''
        study_format = ''
        detail_block = soup.find('div', class_='bannermain1')
        if not detail_block:
            detail_block = soup
        table = detail_block.find('table', class_='td-dist') if detail_block else None
        if table:
            for td in table.find_all('td'):
                span = td.find('span', string=re.compile(r'Стоимость обучения', re.I))
                if span:
                    price_span = td.find('span', string=re.compile(r'\d+[\s\d]*руб', re.I))
                    if price_span:
                        price = price_span.get_text(strip=True)
                    else:
                        price_text = td.get_text(separator=' ', strip=True)
                        match = re.search(r'(\d+[\s\d]*руб\.?)', price_text, re.I)
                        if match:
                            price = match.group(1)
                    break
        if not price:
            for tag in soup.find_all(string=re.compile(r'Стоимость обучения', re.I)):
                parent = tag.find_parent()
                if parent:
                    price_span = parent.find('span', string=re.compile(r'\d+[\s\d]*руб', re.I))
                    if price_span:
                        price = price_span.get_text(strip=True)
                        break
        if not price:
            price_tag = soup.find(string=re.compile(r'\d+[\s\d]*руб', re.I))
            if price_tag:
                price = price_tag.strip()
        if table:
            for td in table.find_all('td'):
                text = td.get_text(separator=' ', strip=True)
                if 'Формат обучения' in text:
                    match = re.search(r'Формат обучения\s*:\s*(.+)', text, re.I)
                    if match:
                        study_format = match.group(1).strip()
                    else:
                        study_format = text
                    break
        if not study_format:
            for tag in soup.find_all(string=re.compile(r'Формат обучения', re.I)):
                parent = tag.find_parent()
                if parent:
                    full_text = parent.get_text(separator=' ', strip=True)
                    match = re.search(r'Формат обучения\s*:\s*(.+)', full_text, re.I)
                    if match:
                        study_format = match.group(1).strip()
                        break
        if not study_format:
            for tag in soup.find_all(string=re.compile(r'очно|заочно|дистанц', re.I)):
                text = tag.strip()
                if len(text) < 30 and any(word in text.lower() for word in ['очно', 'заочно', 'дистанц']):
                    study_format = text
                    break
        return {'price': price, 'format': study_format}

    def parse_main_page():
        soup = get_soup(START_URL)
        if not soup:
            return []
        category_names = [
            "Профессиональная подготовка", "Профессиональная переподготовка",
            "Повышение квалификации", "Дополнительное профессиональное обучение"
        ]
        tab_blocks = soup.select('.first, .second, .third, .fourth')
        all_programs = []
        for idx, block in enumerate(tab_blocks):
            if idx >= len(category_names):
                break
            category = category_names[idx]
            print(f"Обработка категории: {category}")
            items = block.select('.col-xs-6.col-sm-3')
            for item in items:
                link_tag = item.find('a', href=True)
                if not link_tag:
                    continue
                href = link_tag['href']
                full_url = urljoin(BASE_URL, href)
                name_elem = item.select_one('.name_progs')
                name = name_elem.get_text(strip=True) if name_elem else ''
                time_elem = item.select_one('.time_progs')
                hours_text = time_elem.get_text(strip=True) if time_elem else ''
                hours = extract_number(hours_text)
                if not href.startswith('/klientam/perechen-kursov.html'):
                    all_programs.append({
                        'Наименование': name, 'Вид обучения': category,
                        'Количество часов': hours, 'Стоимость': '',
                        'Формат обучения': '', 'Ссылка': full_url
                    })
                    continue
                detail = parse_detail_page(full_url)
                all_programs.append({
                    'Наименование': name, 'Вид обучения': category,
                    'Количество часов': hours, 'Стоимость': detail['price'],
                    'Формат обучения': detail['format'], 'Ссылка': full_url
                })
                time.sleep(0.5)
        return all_programs

    print("=" * 60)
    print("ПАРСЕР ЭНЕРГЕТИК НН (energetik-nn.ru)")
    print("=" * 60)
    programs = parse_main_page()
    if programs:
        df = pd.DataFrame(programs)
        df.to_excel(OUTPUT_FILENAME, index=False)
        print(f"Собрано {len(programs)} программ. Данные сохранены в {OUTPUT_FILENAME}")
    else:
        print("Не удалось найти ни одной программы")

# ========== ПАРСЕР №27: ОСНОВА ==========

def run_parser_osnova():
    """Парсер с сайта pdo-osnova.ru."""
    BASE_URL = 'https://pdo-osnova.ru/'
    HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
    OUTPUT_FILENAME = 'osnova_data.xlsx'
    session = requests.Session()
    session.headers.update(HEADERS)

    def get_soup(url):
        try:
            response = session.get(url, timeout=15, verify=False)
            response.raise_for_status()
            response.encoding = 'utf-8'
            return BeautifulSoup(response.text, 'html.parser')
        except Exception as e:
            print(f'Ошибка при загрузке {url}: {e}')
            return None

    def extract_price(price_text):
        if not price_text:
            return None
        cleaned = re.sub(r'[^\d\s.,]', '', price_text).strip()
        cleaned = re.sub(r'^от\s*', '', cleaned)
        cleaned = cleaned.replace(' ', '').replace(',', '.')
        try:
            return float(cleaned)
        except ValueError:
            return None

    def get_sections():
        soup = get_soup(BASE_URL)
        if not soup:
            known = [
                {'title': 'Повышение квалификации', 'url': urljoin(BASE_URL, 'povyshenie-kvalifikatsii/')},
                {'title': 'Профессиональная переподготовка', 'url': urljoin(BASE_URL, 'professionalnaya-perepodgotovka/')},
                {'title': 'Обучение по рабочим специальностям', 'url': urljoin(BASE_URL, 'obuchenie-rabochim-specialnostyam/')},
                {'title': 'Очное обучение', 'url': urljoin(BASE_URL, 'ochnoe-obuchenie/')},
                {'title': 'Семинары', 'url': urljoin(BASE_URL, 'seminars/')},
                {'title': 'Аттестация персонала', 'url': urljoin(BASE_URL, 'attestaciya/')},
            ]
            return known
        sections = []
        container = soup.find('div', class_='grid-list grid-list--items-2-991')
        if not container:
            containers = soup.find_all('div', class_=re.compile(r'grid-list--items-2'))
            for cont in containers:
                items = cont.find_all('div', class_='banners-img-with-text-list__wrapper')
                if items:
                    container = cont
                    break
        if container:
            items = container.find_all('div', class_='banners-img-with-text-list__wrapper')
            for item in items:
                link_tag = item.find('a', class_='dark_link')
                if link_tag:
                    title = link_tag.get_text(strip=True)
                    href = link_tag.get('href')
                    if href:
                        full_url = urljoin(BASE_URL, href)
                        sections.append({'title': title, 'url': full_url})
        else:
            known_paths = ['povyshenie-kvalifikatsii', 'professionalnaya-perepodgotovka',
                           'obuchenie-rabochim-specialnostyam', 'ochnoe-obuchenie', 'seminars', 'attestaciya']
            for path in known_paths:
                link = soup.find('a', href=re.compile(path))
                if link:
                    title = link.get_text(strip=True)
                    if title:
                        sections.append({'title': title, 'url': urljoin(BASE_URL, link.get('href'))})
            if not sections:
                sections = [
                    {'title': 'Повышение квалификации', 'url': urljoin(BASE_URL, 'povyshenie-kvalifikatsii/')},
                    {'title': 'Профессиональная переподготовка', 'url': urljoin(BASE_URL, 'professionalnaya-perepodgotovka/')},
                    {'title': 'Обучение по рабочим специальностям', 'url': urljoin(BASE_URL, 'obuchenie-rabochim-specialnostyam/')},
                    {'title': 'Очное обучение', 'url': urljoin(BASE_URL, 'ochnoe-obuchenie/')},
                    {'title': 'Семинары', 'url': urljoin(BASE_URL, 'seminars/')},
                    {'title': 'Аттестация персонала', 'url': urljoin(BASE_URL, 'attestaciya/')},
                ]
        return sections

    def get_programs_from_page(soup, section_title):
        programs = []
        items = soup.find_all('div', class_='services-list__wrapper')
        for item in items:
            title_tag = item.find('a', class_='dark_link')
            if not title_tag:
                continue
            title = title_tag.get_text(strip=True)
            href = title_tag.get('href')
            if not href:
                continue
            program_url = urljoin(BASE_URL, href)
            price_tag = item.find('span', class_='price__new-val')
            price_text = price_tag.get_text(strip=True) if price_tag else ''
            price = extract_price(price_text)
            programs.append({'section': section_title, 'name': title, 'url': program_url,
                             'price': price, 'hours': None, 'format': None})
        return programs

    def get_program_details(program_url):
        soup = get_soup(program_url)
        if not soup:
            return {'hours': None, 'format': None}
        hours = None
        format_val = None
        prop_items = soup.find_all('div', itemprop='additionalProperty')
        for prop in prop_items:
            name_span = prop.find('span', itemprop='name')
            if not name_span:
                continue
            name = name_span.get_text(strip=True)
            value_div = prop.find('div', itemprop='value')
            if not value_div:
                continue
            value = value_div.get_text(strip=True)
            if hours is None and 'Количество часов' in name:
                hours = value
            if hours is None and 'Сроки обучения' in name:
                hours = value
            if format_val is None and ('Формат обучения' in name or 'Форма обучения' in name or 'Формат' in name):
                format_val = value
        if hours is None:
            desc_block = soup.find('div', class_='detail-block ordered-block desc')
            if desc_block:
                table = desc_block.find('table')
                if table:
                    ths = table.find_all('th')
                    headers = [th.get_text(strip=True) for th in ths]
                    col_idx = None
                    for i, h in enumerate(headers):
                        if h.lower() in ('длительность', 'объем', 'часы', 'количество часов', 'срок'):
                            col_idx = i
                            break
                    if col_idx is not None:
                        tbody = table.find('tbody')
                        rows = tbody.find_all('tr') if tbody else table.find_all('tr')[1:]
                        for row in rows:
                            cells = row.find_all('td')
                            if len(cells) > col_idx:
                                cell_text = cells[col_idx].get_text(strip=True)
                                if cell_text and ('часов' in cell_text or re.search(r'\d+', cell_text)):
                                    hours = cell_text
                                    break
        if hours is None:
            match = re.search(r'(\d+\s*часов)', soup.get_text())
            if match:
                hours = match.group(1)
        if format_val is None:
            desc_block = soup.find('div', class_='detail-block ordered-block desc')
            if desc_block:
                table = desc_block.find('table')
                if table:
                    ths = table.find_all('th')
                    headers = [th.get_text(strip=True) for th in ths]
                    col_idx = None
                    for i, h in enumerate(headers):
                        if h.lower() in ('формат', 'форма обучения', 'обучение'):
                            col_idx = i
                            break
                    if col_idx is not None:
                        tbody = table.find('tbody')
                        rows = tbody.find_all('tr') if tbody else table.find_all('tr')[1:]
                        for row in rows:
                            cells = row.find_all('td')
                            if len(cells) > col_idx:
                                cell_text = cells[col_idx].get_text(strip=True)
                                if cell_text:
                                    format_val = cell_text
                                    break
        if format_val is None:
            text = soup.get_text()
            format_keywords = ['очное', 'дистанционное', 'заочное', 'очно-заочное', 'вечернее', 'смешанное']
            found = []
            for kw in format_keywords:
                if re.search(r'\b' + kw + r'\b', text, re.IGNORECASE):
                    found.append(kw.capitalize())
            if found:
                format_val = ', '.join(found)
        return {'hours': hours, 'format': format_val}

    def get_all_programs_in_section(section):
        all_programs = []
        url = section['url']
        page_num = 1
        while True:
            if '?' in url:
                paginated_url = url + f'&PAGEN_155={page_num}'
            else:
                paginated_url = url + f'?PAGEN_155={page_num}'
            print(f'Загрузка страницы {page_num} раздела {section["title"]}: {paginated_url}')
            soup = get_soup(paginated_url)
            if not soup:
                break
            programs = get_programs_from_page(soup, section['title'])
            if not programs:
                break
            all_programs.extend(programs)
            pagination = soup.find('div', class_='module-pagination')
            if pagination:
                next_link = pagination.find('a', class_='arrows-pagination__next')
                if not next_link:
                    next_link = pagination.find('a', string=re.compile(r'След\.|>'))
                if next_link:
                    page_num += 1
                    continue
                else:
                    page_links = pagination.find_all('a', class_='module-pagination__item')
                    page_numbers = [int(a.get_text(strip=True)) for a in page_links if a.get_text(strip=True).isdigit()]
                    if page_numbers and max(page_numbers) >= page_num:
                        page_num += 1
                        continue
            break
            time.sleep(1)
        return all_programs

    print("=" * 60)
    print("ПАРСЕР ОСНОВА (pdo-osnova.ru)")
    print("=" * 60)
    sections = get_sections()
    if not sections:
        return
    print(f'Найдено разделов: {len(sections)}')
    all_data = []
    for idx, section in enumerate(sections, 1):
        print(f'\nОбработка раздела {idx}/{len(sections)}: {section["title"]}')
        programs = get_all_programs_in_section(section)
        print(f'Найдено программ: {len(programs)}')
        for prog in programs:
            details = get_program_details(prog['url'])
            prog['hours'] = details['hours']
            prog['format'] = details['format']
            time.sleep(0.5)
        all_data.extend(programs)
    df = pd.DataFrame(all_data, columns=['section', 'name', 'url', 'format', 'hours', 'price'])
    if not df.empty:
        df['price'] = df['price'].astype(float).round(2)
    df = df[['section', 'name', 'url', 'format', 'hours', 'price']]
    df.to_excel(OUTPUT_FILENAME, index=False, engine='openpyxl')
    print(f'Данные сохранены в {OUTPUT_FILENAME}. Всего собрано записей: {len(df)}')

# ========== ПАРСЕР №28: КУБГТУ ==========

def run_parser_kubgtu():
    """Парсер с сайта kubstu.ru."""
    BASE_URL = "https://kubstu.ru/"
    HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/114.0.0.0 Safari/537.36"}
    OUTPUT_FILENAME = 'kubgtu_data.xlsx'

    def clean_dpo_prefix(text):
        if not text:
            return text
        cleaned = re.sub(r'^Программы?\s+ДПО:\s*', '', text, flags=re.IGNORECASE)
        return cleaned.strip()

    def get_soup(url):
        resp = requests.get(url, headers=HEADERS, verify=False, timeout=10)
        resp.raise_for_status()
        resp.encoding = resp.apparent_encoding or 'utf-8'
        return BeautifulSoup(resp.text, "html.parser")

    def parse_program_details(program_url):
        soup = get_soup(program_url)
        article = soup.find("article", class_="mside")
        if not article:
            return {}
        title_tag = article.find("h1")
        title = title_tag.get_text(strip=True) if title_tag else ""
        title = clean_dpo_prefix(title)
        table = article.find("table", class_="sb")
        details = {"Форма обучения": "", "Объем часов": "", "Выдаваемый документ": "", "Стоимость обучения": ""}
        if table:
            for row in table.find_all("tr"):
                th = row.find("th")
                td = row.find("td")
                if not th or not td:
                    continue
                key = th.get_text(strip=True)
                value = td.get_text(strip=True)
                if "Форма обучения" in key:
                    details["Форма обучения"] = value
                elif "Объем часов" in key:
                    details["Объем часов"] = value
                elif "Выдаваемый документ" in key:
                    details["Выдаваемый документ"] = value
                elif "Стоимость" in key:
                    details["Стоимость обучения"] = value
        return {
            "Наименование программы": title, "Ссылка": program_url,
            "Форма обучения": details["Форма обучения"], "Объем часов": details["Объем часов"],
            "Выдаваемый документ": details["Выдаваемый документ"], "Стоимость обучения": details["Стоимость обучения"]
        }

    print("=" * 60)
    print("ПАРСЕР КУБГТУ (kubstu.ru)")
    print("=" * 60)
    all_data = []
    main_soup = get_soup(BASE_URL + "s-22")
    container = main_soup.find("div", class_="dpo-xpelist-box")
    if not container:
        return []
    direction_links = []
    for a in container.find_all("a", href=True):
        href = a["href"]
        if href.startswith("info/xpelist/group-"):
            full_url = urljoin(BASE_URL, href)
            direction_name = a.get_text(strip=True)
            direction_links.append((full_url, direction_name))
    print(f"Найдено направлений: {len(direction_links)}")
    for dir_url, dir_name in direction_links:
        print(f"\nОбрабатываем: {dir_name}")
        try:
            soup = get_soup(dir_url)
        except Exception as e:
            print(f"  Ошибка загрузки {dir_url}: {e}")
            continue
        article = soup.find("article", class_="mside")
        if not article:
            continue
        header = article.find("h1")
        section_name = header.get_text(strip=True) if header else dir_name
        section_name = clean_dpo_prefix(section_name)
        main_ul = article.find("ul", class_="lxm")
        if not main_ul:
            continue
        for li in main_ul.find_all("li", recursive=False):
            dept_tag = li.find("b")
            if not dept_tag:
                continue
            dept_name = dept_tag.get_text(strip=True)
            inner_ul = li.find("ul")
            if not inner_ul:
                continue
            for program_li in inner_ul.find_all("li"):
                a_tag = program_li.find("a", href=True)
                if not a_tag:
                    continue
                program_url = urljoin(BASE_URL, a_tag["href"])
                program_title = a_tag.get_text(strip=True)
                try:
                    details = parse_program_details(program_url)
                except Exception as e:
                    print(f"    Ошибка парсинга {program_url}: {e}")
                    continue
                if not details:
                    continue
                details["Раздел"] = f"{section_name} / {dept_name}"
                if not details["Наименование программы"]:
                    details["Наименование программы"] = program_title
                all_data.append(details)
                time.sleep(0.3)
    if not all_data:
        print("Нет данных для сохранения")
        return
    df = pd.DataFrame(all_data)
    columns_order = ["Наименование программы", "Ссылка", "Раздел", "Форма обучения",
                     "Объем часов", "Выдаваемый документ", "Стоимость обучения"]
    df = df[columns_order]
    df.to_excel(OUTPUT_FILENAME, index=False)
    print(f"Данные сохранены в {OUTPUT_FILENAME}")

# ========== ПАРСЕР №29: УПЦ ==========

def run_parser_upc():
    """Парсер с сайта upc-centr.ru."""
    URL = "https://upc-centr.ru/node/26"
    OUTPUT_FILENAME = 'upc_data.xlsx'

    def get_soup(url):
        headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
        resp = requests.get(url, headers=headers, verify=False, timeout=10)
        resp.encoding = "utf-8"
        return BeautifulSoup(resp.text, "html.parser")

    def clean_text(text):
        if not text:
            return ""
        return " ".join(text.split()).strip()

    def parse_programs(soup):
        table = soup.find("table")
        if not table:
            return []
        rows = table.find_all("tr")
        programs = []
        current_type = "Профессии рабочих"
        type_keywords = {
            "ДОПОЛНИТЕЛЬНЫЕ ПРОФЕССИОНАЛЬНЫЕ ПРОГРАММЫ ПОВЫШЕНИЯ КВАЛИФИКАЦИИ":
                "Дополнительные профессиональные программы повышения квалификации",
            "ДОПОЛНИТЕЛЬНЫЕ ПРОФЕССИОНАЛЬНЫЕ ПРОГРАММЫ ПРОФЕССИОНАЛЬНОЙ ПЕРЕПОДГОТОВКИ":
                "Дополнительные профессиональные программы профессиональной переподготовки"
        }
        for row in rows:
            tds = row.find_all("td")
            if len(tds) == 1 and tds[0].has_attr("colspan"):
                text = clean_text(tds[0].get_text())
                for key, type_name in type_keywords.items():
                    if key in text.upper():
                        current_type = type_name
                        break
                continue
            if row.find("th"):
                continue
            if len(tds) < 9:
                continue
            name_cell = tds[3]
            link = name_cell.find("a")
            name = clean_text(link.get_text()) if link else clean_text(name_cell.get_text())
            if not name:
                continue
            form = clean_text(tds[4].get_text())
            duration = clean_text(tds[5].get_text())
            disciplines = clean_text(tds[6].get_text())
            practice = clean_text(tds[7].get_text())
            cost = clean_text(tds[8].get_text())
            programs.append({
                "Категория": current_type, "Наименование": name, "Форма обучения": form,
                "Нормативный срок": duration, "Учебные дисциплины": disciplines,
                "Практика": practice, "Стоимость": cost
            })
        return programs

    print("=" * 60)
    print("ПАРСЕР УПЦ (upc-centr.ru)")
    print("=" * 60)
    soup = get_soup(URL)
    data = parse_programs(soup)
    if not data:
        print("Данные не найдены")
        return
    print(f"Найдено программ: {len(data)}")
    df = pd.DataFrame(data)
    df.to_excel(OUTPUT_FILENAME, index=False)
    print(f"Данные сохранены в {OUTPUT_FILENAME}")

# ========== ПАРСЕР №30: ЧИПК ==========

def run_parser_chipk():
    """Парсер PDF с сайта chipk.ru."""
    URL = "https://chipk.ru/wp-content/uploads/2026/06/2027_plan-kursov_chf-peipk.pdf"
    OUTPUT_FILENAME = 'chipk_data.xlsx'

    def download_pdf(url):
        headers = {"User-Agent": "Mozilla/5.0"}
        resp = requests.get(url, headers=headers, verify=False, timeout=30)
        resp.raise_for_status()
        return resp.content

    def extract_text_with_pypdf2(pdf_bytes):
        reader = PyPDF2.PdfReader(io.BytesIO(pdf_bytes))
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text

    def parse_courses_improved(text):
        lines = text.split('\n')
        courses = []
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if '«' in line or line.startswith('Ч.'):
                block = line
                j = i + 1
                while j < len(lines) and not lines[j].strip().startswith('Ч.') and lines[j].strip() != '':
                    block += ' ' + lines[j].strip()
                    j += 1
                match = re.search(r'Ч\.\d+\s*[а-я]?\s*«([^»]+)»', block, re.IGNORECASE)
                if not match:
                    match = re.search(r'«([^»]+)»', block)
                if match:
                    name = match.group(1).strip()
                else:
                    name = None
                hours_match = re.search(r'(\d+)\s*(?:час(?:а|ов|)?)', block, re.IGNORECASE)
                hours = int(hours_match.group(1)) if hours_match else None
                cost_match = re.search(r'Стоимость:\s*([\d\s]+)\s*(?:руб\.|рублей|руб)', block, re.IGNORECASE)
                if cost_match:
                    cost_str = cost_match.group(1).replace(' ', '').strip()
                    cost = int(cost_str) if cost_str.isdigit() else cost_str
                else:
                    cost = None
                if name and hours and cost:
                    courses.append({"Наименование": name, "Часы": hours, "Стоимость": cost})
                    i = j
                    continue
            i += 1
        return courses

    print("=" * 60)
    print("ПАРСЕР ЧИПК (chipk.ru)")
    print("=" * 60)
    pdf_bytes = download_pdf(URL)
    text = extract_text_with_pypdf2(pdf_bytes)
    if len(text.strip()) < 100:
        print("PyPDF2 не дал текста, пробуем pdfplumber...")
        try:
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                text = ""
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except ImportError:
            print("pdfplumber не установлен")
            return
    courses = parse_courses_improved(text)
    if not courses:
        print("Курсы не найдены")
        return
    print(f"Найдено курсов: {len(courses)}")
    df = pd.DataFrame(courses)
    df.to_excel(OUTPUT_FILENAME, index=False)
    print(f"Данные сохранены в {OUTPUT_FILENAME}")

# ========== ПАРСЕР №31: ГУЗ ==========

def run_parser_guz():
    """Парсер с сайта guz.ru."""
    URL = "https://www.guz.ru/fakultety/institut-povysheniya-kvalifikatsii/pricesipk.php"
    OUTPUT_FILENAME = 'guz_data.xlsx'

    def get_soup(url):
        headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
        resp = requests.get(url, headers=headers, verify=False, timeout=10)
        resp.encoding = "utf-8"
        return BeautifulSoup(resp.text, "html.parser")

    def clean_text(text):
        if not text:
            return ""
        return " ".join(text.split()).strip()

    def parse_courses(soup):
        table = soup.find("table", {"border": "1"})
        if not table:
            return []
        rows = table.find_all("tr")
        courses = []
        current_type = None
        for row in rows:
            tds = row.find_all("td")
            if len(tds) == 1 and tds[0].has_attr("colspan"):
                text = clean_text(tds[0].get_text())
                if "Профессиональная переподготовка" in text:
                    current_type = "Профессиональная переподготовка"
                elif "Повышение квалификации" in text:
                    current_type = "Повышение квалификации"
                continue
            if row.find("th"):
                continue
            if len(tds) < 5:
                continue
            name_cell = tds[0]
            link = name_cell.find("a")
            name = clean_text(link.get_text()) if link else clean_text(name_cell.get_text())
            if not name:
                continue
            hours = clean_text(tds[1].get_text())
            conditions = clean_text(tds[2].get_text())
            cost = clean_text(tds[3].get_text())
            document = clean_text(tds[4].get_text())
            if current_type is None:
                current_type = "Неизвестно"
            courses.append({
                "Тип программы": current_type, "Наименование": name,
                "Количество часов": hours, "Условия приема": conditions,
                "Стоимость, руб./чел.": cost, "Выдаваемый документ": document
            })
        return courses

    print("=" * 60)
    print("ПАРСЕР ГУЗ (guz.ru)")
    print("=" * 60)
    soup = get_soup(URL)
    data = parse_courses(soup)
    if not data:
        print("Данные не найдены")
        return
    print(f"Найдено курсов: {len(data)}")
    df = pd.DataFrame(data)
    df.to_excel(OUTPUT_FILENAME, index=False)
    print(f"Данные сохранены в {OUTPUT_FILENAME}")

# ========== ОРКЕСТРАТОР ==========

PARSERS = {
    1:  {'name': 'akademik',         'func': run_parser_akademik,         'output_file': 'akademik_data.xlsx'},
    2:  {'name': 'cps',              'func': run_parser_cps,              'output_file': 'cps_data.xlsx'},
    3:  {'name': 'gubkin',           'func': run_parser_gubkin,           'output_file': 'gubkin_data.xlsx'},
    4:  {'name': 'orlovskiy',        'func': run_parser_orlovskiy,        'output_file': 'orlovskiy_data.xlsx'},
    5:  {'name': 'standart_plus',    'func': run_parser_standart_plus,    'output_file': 'standart_plus_data.xlsx'},
    6:  {'name': 'academia',         'func': run_parser_academia,         'output_file': 'academia_data.xlsx'},
    7:  {'name': 'mpei',             'func': run_parser_mpei,             'output_file': 'mpei_data.xlsx'},
    8:  {'name': 'energetik_dpo',    'func': run_parser_energetik_dpo,    'output_file': 'energetik_dpo_data.xlsx'},
    9:  {'name': 'eipk',             'func': run_parser_eipk,             'output_file': 'eipk_data.xlsx'},
    10: {'name': 'ecoips',           'func': run_parser_ecoips,           'output_file': 'ecoips_data.xlsx'},
    11: {'name': 'neftegaz',         'func': run_parser_neftegaz,         'output_file': 'neftegaz_data.xlsx'},
    12: {'name': 'tpu',              'func': run_parser_tpu,              'output_file': 'tpu_data.xlsx'},
    13: {'name': 'az',               'func': run_parser_az,               'output_file': 'az_data.xlsx'},
    14: {'name': 'cppk',             'func': run_parser_cppk,             'output_file': 'cppk_data.xlsx'},
    15: {'name': 'mei_yuga',         'func': run_parser_mei_yuga,         'output_file': 'mei_yuga_data.xlsx'},
    16: {'name': 'upk',              'func': run_parser_upk,              'output_file': 'upk_data.xlsx'},
    17: {'name': 'burovik',          'func': run_parser_burovik,          'output_file': 'burovik_data.xlsx'},
    18: {'name': 'patriot',          'func': run_parser_patriot,          'output_file': 'patriot_data.xlsx'},
    19: {'name': 'chuvash',          'func': run_parser_chuvash,          'output_file': 'chuvash_data.xlsx'},
    20: {'name': 'mrsk_ural',        'func': run_parser_mrsk_ural,        'output_file': 'mrsk_ural_data.xlsx'},
    21: {'name': 'igeu',             'func': run_parser_igeu,             'output_file': 'igeu_data.xlsx'},
    22: {'name': 'cpo',              'func': run_parser_cpo,              'output_file': 'cpo_data.xlsx'},
    23: {'name': 'yurgpu',           'func': run_parser_yurgpu,           'output_file': 'yurgpu_data.xlsx'},
    24: {'name': 'ipo',              'func': run_parser_ipo,              'output_file': 'ipo_data.xlsx'},
    25: {'name': 'lichnostniy_rost', 'func': run_parser_lichnostniy_rost, 'output_file': 'lichnostniy_rost_data.xlsx'},
    26: {'name': 'energetik_nn',     'func': run_parser_energetik_nn,     'output_file': 'energetik_nn_data.xlsx'},
    27: {'name': 'osnova',           'func': run_parser_osnova,           'output_file': 'osnova_data.xlsx'},
    28: {'name': 'kubgtu',           'func': run_parser_kubgtu,           'output_file': 'kubgtu_data.xlsx'},
    29: {'name': 'upc',              'func': run_parser_upc,              'output_file': 'upc_data.xlsx'},
    30: {'name': 'chipk',            'func': run_parser_chipk,            'output_file': 'chipk_data.xlsx'},
    31: {'name': 'guz',              'func': run_parser_guz,              'output_file': 'guz_data.xlsx'},
}

def run_single_parser(parser_id):
    """Запускает один парсер по ID, делает бэкапы и снапшоты."""
    info = PARSERS.get(parser_id)
    if not info:
        logger.error(f"Парсер с ID {parser_id} не найден.")
        return
    name = info['name']
    func = info['func']
    out_file = info['output_file']
    logger.info("=" * 60)
    logger.info(f"ЗАПУСК ПАРСЕРА: {name} (ID={parser_id})")
    logger.info(f"Выходной файл: {out_file}")
    logger.info("=" * 60)
    ensure_dir('backups')
    timestamp = get_timestamp()
    make_backup(out_file, f"{timestamp}_before_{name}")
    start = time.time()
    try:
        func()
        elapsed = time.time() - start
        logger.info(f"Парсер {name} завершился за {elapsed:.2f} сек.")
    except Exception as e:
        logger.error(f"Ошибка при выполнении {name}: {e}", exc_info=True)
    finally:
        make_backup(out_file, f"{timestamp}_after_{name}")
        save_snapshot(out_file, name)
    logger.info(f"Завершена обработка парсера {name}\n")

def print_menu():
    print("\n" + "=" * 60)
    print("ОРКЕСТРАТОР ПАРСЕРОВ ОБРАЗОВАТЕЛЬНЫХ ОРГАНИЗАЦИЙ")
    print("=" * 60)
    for pid, info in PARSERS.items():
        print(f"  {pid:>2}. {info['name']:<20} -> {info['output_file']}")
    print("  all - запустить все парсеры")
    print("  q   - выход")
    print("=" * 60)

def parse_choice(user_input):
    user_input = user_input.strip().lower()
    if user_input == 'q':
        return None
    if user_input == 'all':
        return list(PARSERS.keys())
    ids = []
    parts = user_input.replace(' ', '').split(',')
    for p in parts:
        if '-' in p:
            try:
                start, end = map(int, p.split('-'))
                ids.extend(range(start, end + 1))
                continue
            except Exception:
                pass
        if p.isdigit():
            num = int(p)
            if num in PARSERS:
                ids.append(num)
            else:
                logger.warning(f"Неизвестный номер: {num}")
        else:
            logger.warning(f"Некорректный ввод: {p}")
    return list(set(ids))

def main():
    logger.info("=" * 60)
    logger.info("ЗАПУСК ОРКЕСТРАТОРА (31 парсер в одном файле)")
    logger.info(f"Время старта: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    logger.info("=" * 60)
    while True:
        print_menu()
        choice = input("Введите номера парсеров (через запятую/диапазон), 'all' или 'q': ").strip()
        if choice.lower() == 'q':
            logger.info("Выход по запросу пользователя.")
            break
        ids = parse_choice(choice)
        if ids is None:
            continue
        if not ids:
            logger.warning("Не выбрано ни одного парсера.")
            continue
        logger.info(f"Выбраны парсеры: {', '.join(str(i) for i in sorted(ids))}")
        for pid in sorted(ids):
            run_single_parser(pid)
        logger.info("Все выбранные парсеры выполнены.\n")
    logger.info("Оркестратор завершил работу.")

if __name__ == "__main__":
    main()
    